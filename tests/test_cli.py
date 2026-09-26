"""hate 리뷰(루프 1) first_nail 회귀: 읽을 수 없는 양식을 '성공'으로 내지 않는다."""
import json

import openpyxl
import pytest

from danburn.cli import main


def _xlsx(path, title, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_unreadable_format_fails_loudly_and_writes_nothing(tmp_path, capsys):
    src = tmp_path / "synth.xlsx"
    _xlsx(src, "내역", [["품명", "규격", "단위", "수량"], ["레미콘", "25-24-15", "m3", 500], ["이형철근", "SD400 D13", "ton", 30]])
    out = tmp_path / "o.hwpx"
    assert main(["build", "--boq", str(src), "--out", str(out)]) == 2
    assert not out.exists()
    assert "내역" in capsys.readouterr().err       # 어떤 시트가 있었는지 알려 준다


def test_material_with_no_rows_is_warned(tmp_path, capsys):
    src = tmp_path / "supplied.xlsx"
    _xlsx(src, "지급(건)", [["품명", "규격", "단위", "수량"], ["레미콘", "25-24-15", "M3", 500]])
    out = tmp_path / "o.hwpx"
    assert main(["build", "--boq", str(src), "--out", str(out), "--offline"]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert any("rebar" in w or "철근" in w for w in summary["warnings"])


def test_ambiguous_block_is_rejected(tmp_path, capsys):
    src = tmp_path / "s.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "지급(건)"
    ws.append(["품명", "규격", "단위", "A동", "", "PA동"])
    ws.append(["", "", "", "수량", "금액", "수량"])
    ws.append(["레미콘", "25-24-15", "M3", 100, "", 50])
    wb.save(src)
    assert main(["inspect", "--boq", str(src)]) == 0
    info = json.loads(capsys.readouterr().out)
    assert info["supported"] and set(info["blocks"]) == {"A동", "PA동"}
    assert main(["build", "--boq", str(src), "--block", "A", "--out", str(tmp_path / "o.hwpx")]) == 2
    assert "여러 블록" in capsys.readouterr().err


def test_civil_scope_and_missing_common_specs(tmp_path, capsys):
    src = tmp_path / "s.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "지급(토)"
    ws.append(["품명", "규격", "단위", "A동", "", "B동"])
    ws.append(["", "", "", "수량", "금액", "수량"])
    ws.append(["레미콘", "25-21-15", "M3", 100, "", 50])
    ws2 = wb.create_sheet("지급(건)")
    ws2.append(["품명", "규격", "단위", "A동", "", "B동"])
    ws2.append(["", "", "", "수량", "금액", "수량"])
    ws2.append(["레미콘", "25-24-15", "M3", 30, "", 20])
    wb.save(src)
    out = tmp_path / "o.hwpx"
    assert main(["build", "--boq", str(src), "--block", "B동", "--civil-scope", "공구",
                 "--add-spec", "건축:rebar:SD400 D10", "--out", str(out)]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert "건축:rebar:SD400 D10" in summary["missing_common_specs"]
    rows = json.loads(out.with_suffix(".json").read_text(encoding="utf-8"))
    civil = {r["qty"] for r in rows if r["discipline"] == "토목"}
    arch = {r["qty"] for r in rows if r["discipline"] == "건축" and r["material"] == "ready_mixed_concrete"}
    assert civil == {150.0} and arch == {20.0}                    # 토목은 두 블록 합계, 건축은 B동만
    added = [r for r in rows if r["material"] == "rebar"]
    assert added and "누락 추정" in added[0]["note"]


def test_plan_command_builds_whole_document(tmp_path, capsys):
    from pathlib import Path
    src = tmp_path / "s.xlsx"
    _xlsx(src, "지급(건)", [["품명", "규격", "단위", "수량"], ["레미콘", "25-24-15", "M3", 500]])
    project = Path(__file__).resolve().parents[1] / "src" / "danburn" / "data" / "templates" / "project.example.yaml"
    out = tmp_path / "plan.hwpx"
    assert main(["plan", "--boq", str(src), "--project", str(project), "--revision", "0",
                 "--date", "2026. 01. 05.", "--out", str(out)]) == 0
    assert out.exists() and out.stat().st_size > 20000
    bad = tmp_path / "bad.hwpx"                              # 예시 Rev.0 일자와 다른 날짜 → 모순, 파일 없음
    assert main(["plan", "--boq", str(src), "--project", str(project), "--revision", "0",
                 "--date", "2026. 09. 26.", "--out", str(bad)]) == 2
    assert not bad.exists() and "Rev" in capsys.readouterr().err


def test_uncovered_materials_are_reported(tmp_path, capsys):
    """hate 루프 3 first_nail: 규칙 없는 자재가 조용히 빠지지 않는다(경고 + 본문 표시)."""
    src = tmp_path / "b.xlsx"
    _xlsx(src, "지급(건)", [["품명", "규격", "단위", "수량"],
                           ["레미콘", "25-24-15", "M3", 500], ["아스팔트콘크리트", "#78", "TON", 300],
                           ["시멘트벽돌", "190*90*57", "EA", 10000], ["복층유리", "24mm", "M2", 800]])
    import shutil
    from pathlib import Path
    only = tmp_path / "rules"                      # 규칙이 늘어나도 흔들리지 않게 레미콘 규칙만 둔다
    only.mkdir()
    shutil.copy(Path(__file__).resolve().parents[1] / "src" / "danburn" / "data" / "rules" / "ready_mixed_concrete.yaml", only)
    out = tmp_path / "o.hwpx"
    assert main(["build", "--boq", str(src), "--rules", str(only), "--out", str(out)]) == 0
    summary = json.loads(capsys.readouterr().out)
    missing = " ".join(w for w in summary["warnings"])
    for word in ("아스팔트", "벽돌", "유리"):
        assert word in missing, word
    assert summary["uncovered_materials"]


def test_needs_confirmation_collects_row_notes(tmp_path, capsys):
    src = tmp_path / "s.xlsx"
    _xlsx(src, "지급(건)", [["품명", "규격", "단위", "수량"], ["레미콘", "25-24-15", "M3", 20]])   # 40㎥ 미만 → 생략 가능
    assert main(["build", "--boq", str(src), "--out", str(tmp_path / "o.hwpx")]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert any(x["note"] == "생략 가능" and x["rows"] >= 1 for x in summary["needs_confirmation"])

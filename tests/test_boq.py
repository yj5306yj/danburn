"""지급자재 내역서 파서 — 합성 xlsx만 쓴다(실자료 금지)."""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from danburn.boq import normalize_concrete_spec, normalize_rebar_spec, normalize_unit, read_boq


def _supply_sheet(ws, rows):
    """실무 양식을 흉내 낸 3단 머리: 품명/규격/단위 · 블록명 · 수량/금액."""
    ws["A1"], ws["E1"] = "내역\n구분", "합성 내역서 [지급]"
    ws["A2"], ws["E2"] = "1.목차", "◈ 공사명 : 합성공사"
    for r in (3, 4, 5):
        ws.cell(r, 1, "1.목차")
    ws["E3"], ws["F3"], ws["G3"], ws["H3"] = "품명", "규격", "단위", "설 계 금 액"
    for col in "EFG":
        ws.merge_cells(f"{col}3:{col}5")
    ws.merge_cells("H3:P3")
    ws["H4"], ws["J4"], ws["M4"], ws["O4"] = "단 가", "블록가", "블록나", "합 계"
    ws.merge_cells("H4:I4")
    ws.merge_cells("J4:L4")
    ws.merge_cells("M4:N4")
    ws.merge_cells("O4:P4")
    ws["J5"], ws["K5"], ws["M5"], ws["N5"], ws["O5"], ws["P5"] = "수량", "금 액", "수 량", "금 액", "수량", "금 액"
    for i, row in enumerate(rows, start=6):
        first, name, spec, unit, qa, qb = row
        ws.cell(i, 1, first)
        ws.cell(i, 5, name)
        ws.cell(i, 6, spec)
        ws.cell(i, 7, unit)
        ws.cell(i, 10, qa)
        ws.cell(i, 13, qb)
        total = (qa or 0) + (qb or 0) if isinstance(qa, (int, float)) and isinstance(qb, (int, float)) else None
        ws.cell(i, 15, total)
        ws.cell(i, 11, 1000)   # 금액 열은 무시돼야 함


@pytest.fixture
def boq_file(tmp_path):
    wb = Workbook()
    wb.active.title = "표지"
    wb.active["A1"] = "합성 표지"
    _supply_sheet(wb.create_sheet("지급(건)"), [
        ("1.목차", "1. 건축공사", None, None, None, None),
        ("1.구분", "1-1. 가동", None, None, None, None),
        (None, "레미콘", "25-18-8", "M3", 100, 50),
        (None, "레미콘", "25-18-80", "m3", 20, None),
        (None, "레미콘", "25 - 24 - 150", "㎥", 0, 30),
        (None, "레미콘", "25-24-15", "M3", "1,200", 0),
        (None, "콘크리트벽돌", "190X90X57", "매", 500.5, 0),
        ("1.소계", "소계", None, None, 120, 80),
        (None, "소계", None, None, 1, 1),
    ])
    _supply_sheet(wb.create_sheet("지급(토)"), [
        (None, "철근", "SD400 D13", "TON", 1.5, 2.25),
        (None, "철근", "SD500 D16", "톤", None, 3),
    ])
    ws = wb.create_sheet("내역(건)")        # 이번 범위 밖 — 건너뛴다
    ws["A1"] = "수량"
    path = tmp_path / "synthetic.xlsx"
    wb.save(path)
    return path


def test_reads_blocks_and_skips_markers_totals_zeros(boq_file):
    lines = read_boq(boq_file)
    got = {(l.discipline, l.sheet, l.name, l.spec, l.unit, l.qty, l.block) for l in lines}
    assert got == {
        ("건축", "지급(건)", "레미콘", "25-18-8", "M3", 100.0, "블록가"),
        ("건축", "지급(건)", "레미콘", "25-18-8", "M3", 50.0, "블록나"),
        ("건축", "지급(건)", "레미콘", "25-18-80", "m3", 20.0, "블록가"),
        ("건축", "지급(건)", "레미콘", "25 - 24 - 150", "㎥", 30.0, "블록나"),
        ("건축", "지급(건)", "레미콘", "25-24-15", "M3", 1200.0, "블록가"),
        ("건축", "지급(건)", "콘크리트벽돌", "190X90X57", "매", 500.5, "블록가"),
        ("토목", "지급(토)", "철근", "SD400 D13", "TON", 1.5, "블록가"),
        ("토목", "지급(토)", "철근", "SD400 D13", "TON", 2.25, "블록나"),
        ("토목", "지급(토)", "철근", "SD500 D16", "톤", 3.0, "블록나"),
    }
    assert all(l.supply == "지급" for l in lines)
    assert all(l.block != "합 계" for l in lines)


def test_supply_lines_get_section(boq_file):
    lines = read_boq(boq_file)
    assert {l.section for l in lines if l.sheet == "지급(건)"} == {"가동"}
    assert {l.section for l in lines if l.sheet == "지급(토)"} == {""}


def _detail_sheet(ws, rows):
    """사급 내역 양식: 금액 묶음 2개(설계·도급) × 변경 전/후/증감, 블록 열 없음, 번호 제목 행."""
    ws["A1"], ws["E1"] = "내역\n구분", "합성 공사 내역서"
    for r in (2, 3, 4, 5):
        ws.cell(r, 1, "1.목차")
    ws["E3"], ws["F3"], ws["G3"] = "품명", "규격", "단위"
    for col in "EFG":
        ws.merge_cells(f"{col}3:{col}5")
    ws["H3"], ws["N3"] = "설 계 금 액", "도 급 금 액"
    ws.merge_cells("H3:M3")
    ws.merge_cells("N3:S3")
    labels = ["변 경 전 [A]", "변 경 후 [B]", "증 감 [B-A]", "변 경 전 [C]", "변 경 후 [D]", "증 감 [D-C]"]
    for i, label in enumerate(labels):
        col = 8 + 2 * i
        ws.cell(4, col, label)
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        ws.cell(5, col, "수량")
        ws.cell(5, col + 1, "금액")
    for i, (first, name, spec, unit, before, after) in enumerate(rows, start=6):
        ws.cell(i, 1, first)
        cell = ws.cell(i, 5, name)
        if name.startswith("="):
            cell.data_type = "s"           # "=== 제목 ===" 은 실자료에서 문자열(수식 아님)
        ws.cell(i, 6, spec)
        ws.cell(i, 7, unit)
        if before is not None:
            ws.cell(i, 8, before + 1000)   # 설계 묶음은 쓰지 않는다(다른 값으로 구분)
            ws.cell(i, 10, after + 1000)
            ws.cell(i, 14, before)
            ws.cell(i, 16, after)
            ws.cell(i, 18, after - before)


@pytest.fixture
def detail_file(tmp_path):
    wb = Workbook()
    wb.active.title = "원가(건)"                 # 범위 밖 시트
    _detail_sheet(wb.create_sheet("내역(건)가"), [
        ("1.목차", "1. 합성건축", None, None, None, None),
        ("1.구분", "1-1. 가동", None, None, None, None),
        ("1.구분", "1-1 01. 상부공사", None, None, None, None),
        ("1.구분", "1-1 0101. 철근콘크리트공사", None, None, None, None),
        ("일반", "이형봉강(SD500, 현장도착도)", "H-13", "TON", 10, 12),
        ("별산자재", "철근 콘크리트타설(펌프카)", "S15cm, 보통", "M3", 300, 300),
        ("1.소계", "소계", None, None, 310, 312),
        ("1.구분", "1-1 02. 기초공사", None, None, None, None),
        ("1.구분", "1-1 0201. 철근콘크리트공사", None, None, None, None),
        ("일반", "이형봉강(SD600, 현장도착도)", "H-22", "TON", 5, 5),
        ("일반", "철근 공장가공 및 조립", "보통", "TON", 5, 5),
        ("일반", "C급 콘크리트 치기(레미콘별도)", "25-18-8", "M3", 0, 0),
        ("1.구분", "1-2. 나동", None, None, None, None),
        ("1.구분", "1-2 0101. 철근콘크리트공사", None, None, None, None),
        ("일반", "이형봉강(SD400, 현장도착도)", "D-10", "TON", 2, 0),
    ])
    _detail_sheet(wb.create_sheet("내(토)가"), [
        ("1.구분", "1. 옹벽공사", None, None, None, None),
        ("1.구분", "=== 역L형 옹벽 ===", None, None, None, None),
        ("일반", "철근콘크리트타설/펌프카", "보통", "M3", 40, 40),
        ("1.구분", "=== 이하 공통자재 ===", None, None, None, None),
        ("일반", "레미콘", "25-21-15", "M3", 7, 9),
        ("1.구분", "2. 배수공사", None, None, None, None),
        ("일반", "우수관", "D-450", "M", 3, 3),
    ])
    path = tmp_path / "detail.xlsx"
    wb.save(path)
    return path


def test_detail_sheet_reads_contract_after_quantities_with_sections(detail_file):
    lines = read_boq(detail_file)
    got = [(l.sheet, l.name, l.spec, l.qty, l.block, l.supply, l.section) for l in lines]
    assert got == [
        ("내역(건)가", "이형봉강(SD500, 현장도착도)", "H-13", 12.0, "가", "사급", "가동 > 상부공사 > 철근콘크리트공사"),
        ("내역(건)가", "철근 콘크리트타설(펌프카)", "S15cm, 보통", 300.0, "가", "사급", "가동 > 상부공사 > 철근콘크리트공사"),
        ("내역(건)가", "이형봉강(SD600, 현장도착도)", "H-22", 5.0, "가", "사급", "가동 > 기초공사 > 철근콘크리트공사"),
        ("내역(건)가", "철근 공장가공 및 조립", "보통", 5.0, "가", "사급", "가동 > 기초공사 > 철근콘크리트공사"),
        ("내(토)가", "철근콘크리트타설/펌프카", "보통", 40.0, "가", "사급", "옹벽공사 > 역L형 옹벽"),
        ("내(토)가", "레미콘", "25-21-15", 9.0, "가", "사급", "옹벽공사 > 이하 공통자재"),
        ("내(토)가", "우수관", "D-450", 3.0, "가", "사급", "배수공사"),
    ]
    assert {l.discipline for l in lines if l.sheet == "내(토)가"} == {"토목"}


def test_rebar_material_lines_normalize(detail_file):
    rebar = sorted((normalize_rebar_spec(l.name, l.spec), l.qty) for l in read_boq(detail_file)
                   if normalize_rebar_spec(l.name, l.spec))
    assert rebar == [("SD500 D13", 12.0), ("SD600 D22", 5.0)]


@pytest.mark.parametrize("name, spec, expected", [
    ("이형봉강(SD500, 현장도착도)", "H-13", "SD500 D13"),
    ("이형봉강(SD300, 현장도착도)", "D-10", "SD300 D10"),
    ("철근", "SD400 D13", "SD400 D13"),
    ("이형철근", "SD600, HD22", "SD600 D22"),
    ("철근 공장가공 및 조립", "보통", None),
    ("철근시공도 제작", "", None),
    ("오.배수용 PVC 파이프", "D50 MM(VG2)", None),
])
def test_normalize_rebar_spec(name, spec, expected):
    assert normalize_rebar_spec(name, spec) == expected


def test_row_numbers_point_to_source(boq_file):
    first = next(l for l in read_boq(boq_file) if l.spec == "25-18-8" and l.block == "블록가")
    assert (first.sheet, first.row) == ("지급(건)", 8)


def test_change_statement_uses_after_group(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "지급(기)"
    ws["A2"] = "1.목차"
    ws["C3"], ws["D3"], ws["E3"] = "품명", "규격", "단위"
    ws["F3"], ws["H3"], ws["J3"] = "변 경 전 [A]", "변 경 후 [B]", "증 감 [B-A]"
    for a, b in (("F3", "G3"), ("H3", "I3"), ("J3", "K3")):
        ws.merge_cells(f"{a}:{b}")
    for col, blk in zip("FGHIJK", ["블록가", "합 계"] * 3):
        ws[f"{col}4"] = blk
        ws[f"{col}5"] = "수량"
    ws["C6"], ws["D6"], ws["E6"] = "배관", "D-100", "M"
    ws["F6"], ws["G6"], ws["H6"], ws["I6"], ws["J6"], ws["K6"] = 10, 10, 12, 12, 2, 2
    path = tmp_path / "change.xlsx"
    wb.save(path)
    lines = read_boq(path)
    assert [(l.discipline, l.block, l.qty, l.unit) for l in lines] == [("기계", "블록가", 12.0, "M")]


def test_change_form_without_after_values_falls_back_to_before(tmp_path):
    """블록 열 없이 변경 전/후/증감만 있고 '변경 후'가 빈 양식 → '변경 전' 수량, 블록명은 빈 문자열."""
    wb = Workbook()
    ws = wb.active
    ws.title = "지급(토)"
    ws["E3"], ws["F3"], ws["G3"], ws["H3"] = "품명", "규격", "단위", "설 계 금 액"
    ws.merge_cells("H3:M3")
    ws["H4"], ws["J4"], ws["L4"] = "변 경 전 [A]", "변 경 후 [B]", "증 감 [B-A]"
    for a, b in (("H4", "I4"), ("J4", "K4"), ("L4", "M4")):
        ws.merge_cells(f"{a}:{b}")
    ws["H5"], ws["J5"], ws["L5"] = "수량", "수량", "수량"
    ws["E6"], ws["F6"], ws["G6"], ws["H6"] = "레미콘", "25-21-15", "m3", 40
    ws["A7"], ws["E7"], ws["J7"] = "1.소계", "소계", 99        # 표시 행의 값은 판단에 쓰지 않는다
    path = tmp_path / "before_only.xlsx"
    wb.save(path)
    lines = read_boq(path)
    assert [(l.discipline, l.block, l.spec, l.qty) for l in lines] == [("토목", "", "25-21-15", 40.0)]


def test_rejects_lock_file(tmp_path):
    lock = tmp_path / "~$synthetic.xlsx"
    lock.write_bytes(b"")
    with pytest.raises(ValueError):
        read_boq(lock)


@pytest.mark.parametrize("raw, expected", [
    ("25-18-8", "25-18-80"),
    ("25-18-80", "25-18-80"),
    ("25 - 18 - 150", "25-18-150"),
    ("25-24-15", "25-24-150"),
    ("25-21-12 (별산)", "25-21-120"),
    ("25-24-8,TC기초", "25-24-80"),
    ("SD400 D13", None),
    ("190X90X57", None),
    ("", None),
])
def test_normalize_concrete_spec(raw, expected):
    assert normalize_concrete_spec(raw) == expected


@pytest.mark.parametrize("raw, expected", [
    ("M3", "m3"), ("m3", "m3"), ("㎥", "m3"),
    ("TON", "ton"), ("톤", "ton"), ("t", "ton"),
    ("M2", "m2"), ("㎡", "m2"),
    ("매", "매"), ("EA", "ea"), (" 개소 ", "개소"),
])
def test_normalize_unit(raw, expected):
    assert normalize_unit(raw) == expected


def test_sheet_tail_block_is_mapped_to_column_block_label(tmp_path):
    """블록별 시트 꼬리('3')가 같은 파일의 블록 열 이름('A동-3')과 하나로 맞으면 그 이름을 쓴다."""
    import openpyxl
    from danburn.boq import read_boq
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "지급(건)"
    ws.append(["품명", "규격", "단위", "A동-2", "", "A동-3"])
    ws.append(["", "", "", "수량", "금액", "수량"])
    ws.append(["레미콘", "25-24-15", "M3", 100, "", 50])
    ws2 = wb.create_sheet("내(건)3")
    ws2.append(["품명", "규격", "단위", "수량"])
    ws2.append(["이형봉강(SD400)", "D-13", "TON", 7])
    p = tmp_path / "b.xlsx"
    wb.save(p)
    blocks = {(ln.sheet, ln.block) for ln in read_boq(p)}
    assert ("내(건)3", "A동-3") in blocks

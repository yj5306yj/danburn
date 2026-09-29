"""품질시험계획서 엑셀 단독본(L14-E) — 합성 행만 쓴다. 다시 읽어 머리·병합·인쇄 설정·숫자형을 본다."""
import openpyxl
import pytest

from danburn.hwpx_out import HEAD, merge_plan
from danburn.model import PlanRow
from danburn.testplan import SCAN, STAFF_PROOFS
from danburn.testplan_xlsx import PLAN_BODY_ROW, PLAN_HEAD_ROW, build_workbook, save_workbook

DATE = "2026. 01. 05."


def _row(disc, work, item, test, qty, unit, n_site=0, n_ext=0, ks="", method=""):
    r = PlanRow(disc, work, item, test, qty, unit, "합성 빈도", "합성 근거", n_site, n_ext, ks, material=item)
    r.method = method                       # PlanRow.method 필드가 생기기 전에도 같은 방식으로 읽는다(getattr)
    return r


def _rows():
    return [_row("건축", "철근콘크리트공사", "레미콘(25-24-150)", "슬럼프", 480, "㎥", 4, method="KS F 2402"),
            _row("건축", "철근콘크리트공사", "레미콘(25-24-150)", "공기량", 480, "㎥", 4, method="KS F 2421"),
            _row("건축", "방수공사", "합성 방수시트", "인장강도", 1200.5, "㎡", 0, 1, method="KS F 4917"),
            _row("건축", "방수공사", "합성 방수시트", "겉모양", 1200.5, "㎡", 0, 0, "◎"),
            _row("토목", "토공사", "합성 성토재", "다짐", 3000, "㎥", 2, method="KS F 2312")]


def _project():
    return {"공사명": "가상 합성 공동주택 신축공사", "회사명": "합성건설", "문서번호": "TP-EX-01", "제정일자": DATE,
            "개정이력": [{"개정": 0, "일자": DATE, "장": ["전체"], "사유": "최초 제정"}],
            "시험장비": [{"시험기구": "합성 압축시험기", "규격": "1,000kN", "단위": "대", "수량": "1", "제작사": "합성",
                       "기기번호": "X-1", "교정주기": "1회/년", "비고": ""}],
            "품질관리자": [{"직무": "품질관리자", "성명": "합성 갑", "등급": "고급", "배치기간": "착공~준공", "비고": ""}],
            "조직": [{"직무": "현장대리인", "성명": "합성 을", "자격": "건축기사", "담당": "총괄"}],
            "시험실": "50㎡ 이상(합성)"}


@pytest.fixture
def wb(tmp_path):
    out = save_workbook(_rows(), _project(), tmp_path / "t.xlsx", basis_version="합성 기준판", revision=0, date=DATE,
                        unlisted=[{"kind": "uncovered", "label": "합성 자재", "basis": "별표2", "lines": 3}])
    return openpyxl.load_workbook(out)


def test_sheets_follow_four_parts(wb):
    assert wb.sheetnames == ["표지", "1.공사개요", "2.시험계획-건축", "2.시험계획-토목", "2.미작성 자재(확인)",
                             "3.시험장비·교정", "3.시험실", "4.배치계획"]
    cover = " ".join(str(c.value) for row in wb["표지"].iter_rows() for c in row if c.value)
    for part in ("개요", "품질시험 및 검사계획", "품질시험 시설", "품질관리자 배치계획", "품질시험계획서", "Rev.0"):
        assert part in cover


def test_plan_sheet_header_same_as_hwpx(wb):
    ws = wb["2.시험계획-건축"]
    top = [ws.cell(PLAN_HEAD_ROW, c).value for c in range(1, len(HEAD) + 1)]
    low = [ws.cell(PLAN_HEAD_ROW + 1, c).value for c in range(1, len(HEAD) + 1)]
    assert top[:7] == HEAD[:7] and top[8] == "계획시험횟수" and top[11] == "비고"
    assert low[7:11] == ["산출근거", "현장", "의뢰", "KS"]
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert {"A2:A3", "D2:D3", "G2:H2", "I2:K2", "L2:L3"} <= merged           # 2줄 머리(hwpx HEAD_CELLS 와 같음)


def test_body_merges_match_merge_plan_and_numbers_are_numeric(wb):
    ws = wb["2.시험계획-건축"]
    rows = [r for r in _rows() if r.discipline == "건축"]
    merged = {str(m) for m in ws.merged_cells.ranges}
    for c, runs in merge_plan(rows).items():
        for s, e in runs:
            if e > s:
                col = openpyxl.utils.get_column_letter(c + 1)
                assert f"{col}{PLAN_BODY_ROW + s}:{col}{PLAN_BODY_ROW + e}" in merged
    assert ws.cell(PLAN_BODY_ROW, 4).value == "KS F 2402"                    # 시험방법 칸
    qty = ws.cell(PLAN_BODY_ROW + 2, 5)
    assert qty.value == 1200.5 and isinstance(qty.value, float) and "#,##0" in qty.number_format
    assert ws.cell(PLAN_BODY_ROW, 9).value == 4 and ws.cell(PLAN_BODY_ROW + 2, 10).value == 1   # 현장·의뢰 = 정수
    assert ws.cell(PLAN_BODY_ROW + 3, 11).value == "◎"
    for row in ws.iter_rows():
        for c in row:
            assert not (isinstance(c.value, str) and c.value.startswith("="))  # 수식 없음


def test_print_setup_landscape_repeat_fit_and_breaks(wb):
    ws = wb["2.시험계획-건축"]
    ps = ws.page_setup
    assert ps.orientation == "landscape" and int(ps.paperSize) == 9
    assert int(ps.fitToWidth) == 1 and int(ps.fitToHeight) == 0 and ws.sheet_properties.pageSetUpPr.fitToPage
    assert ws.print_title_rows == "$2:$3"
    assert [b.id for b in ws.row_breaks.brk] == [PLAN_BODY_ROW + 2 - 1]      # 방수공사(공종 경계) 앞
    assert ws.print_area.endswith(f"$L${PLAN_BODY_ROW + 3}")
    for name in ("표지", "3.시험실", "4.배치계획"):
        assert wb[name].page_setup.orientation == "portrait"


def test_scan_placeholders_and_project_tables(wb):
    room = " ".join(str(c.value) for row in wb["3.시험실"].iter_rows() for c in row if c.value)
    assert f"{SCAN} 시험실 배치평면도" in room and "50㎡ 이상(합성)" in room
    staff = " ".join(str(c.value) for row in wb["4.배치계획"].iter_rows() for c in row if c.value)
    assert "합성 갑" in staff and all(f"{SCAN} {p}" in staff for p in STAFF_PROOFS)
    eq = wb["3.시험장비·교정"]
    assert [c.value for c in eq[3]][:4] == ["번호", "시험기구", "규격", "단위"] and eq.cell(4, 2).value == "합성 압축시험기"
    assert "X-1" in [c.value for c in eq[4]]


def test_empty_project_still_builds_with_blanks():
    wb = build_workbook(_rows(), {}, date=DATE)
    ov = " ".join(str(c.value) for row in wb["1.공사개요"].iter_rows() for c in row if c.value)
    assert "(작성 필요)" in ov
    assert wb["3.시험장비·교정"].max_row >= 3 + 5                               # 빈 줄 5줄(사용자가 채운다)


def test_revision_mismatch_raises():
    p = _project()
    with pytest.raises(ValueError):
        build_workbook(_rows(), p, revision=0, date="2026. 02. 01.")

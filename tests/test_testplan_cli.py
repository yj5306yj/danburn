"""danburn test-plan(L14-G): start 산출 폴더의 project.yaml 하나로 품질시험계획서 한글·엑셀·JSON — 합성 내역서만 쓴다.

수록본(start 가 만든 품질관리계획서의 8.11)과 단독본(test-plan)이 같은 compute() 결과를 쓰는지 본다."""
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest
import yaml

from conftest import VALIDATE, run_hwpx_validate
from test_start import _answers, _lines, boq  # noqa: F401  (boq 는 pytest fixture — 질문 순서는 start 테스트가 관리)

DANBURN = Path(sys.executable).parent / "danburn"


def _run(args, home, stdin=None):
    env = {**os.environ, "DANBURN_HOME": str(home)}
    return subprocess.run([str(DANBURN), *args], input=stdin, capture_output=True, text=True, encoding="utf-8",
                          env=env, timeout=300)


@pytest.fixture(scope="module")
def started(boq, tmp_path_factory):  # noqa: F811
    """start --answers(비대화형)로 만든 산출 폴더(project.yaml·품질관리계획서.hwpx·.json)."""
    d = tmp_path_factory.mktemp("tp")
    ans = d / "ans.yaml"
    ans.write_text(yaml.safe_dump(_answers(boq), allow_unicode=True), encoding="utf-8")
    out = d / "out"
    r = _run(["start", "--answers", str(ans), "--yes", "--out-dir", str(out), "--offline"], d / "home")
    assert r.returncode == 0, r.stdout + r.stderr
    assert "danburn test-plan --project" in r.stdout                   # 끝 안내 한 줄(질문은 늘리지 않음)
    return out


@pytest.fixture(scope="module")
def tp(started, tmp_path_factory):
    r = _run(["test-plan", "--project", str(started / "project.yaml"), "--offline", "--json"], tmp_path_factory.mktemp("h"))
    assert r.returncode == 0, r.stdout + r.stderr
    return json.loads(r.stdout)


def test_test_plan_writes_hwpx_xlsx_json_next_to_project(started, tp):
    for ext in ("hwpx", "xlsx", "json"):
        assert (started / f"품질시험계획서.{ext}").stat().st_size > 0
    assert tp["hwpx"] == str(started / "품질시험계획서.hwpx") and tp["xlsx"] == str(started / "품질시험계획서.xlsx")
    assert tp["test_plan"]["scan_attachments"][0] == "시험실 배치평면도"
    assert tp["test_plan"]["block"] == "나동" and tp["basis_check"]["status"] == "unknown"   # --offline
    assert not list(started.glob(".*.tmp*"))                              # 임시 파일이 남지 않는다


def test_rows_equal_plan_811(started, tp):
    """수록본(품질관리계획서.json = plan 의 8.11 행)과 단독본 JSON 이 같고, 행 수가 요약·엑셀과 맞는다."""
    plan = json.loads((started / "품질관리계획서.json").read_text(encoding="utf-8"))
    single = json.loads((started / "품질시험계획서.json").read_text(encoding="utf-8"))
    assert single == plan and len(single) == tp["test_plan"]["rows"] == tp["plan_rows"] > 0
    wb = openpyxl.load_workbook(started / "품질시험계획서.xlsx")
    body = sum(ws.max_row - 3 for ws in wb.worksheets if ws.title.startswith("2.시험계획-"))   # 제목 1 + 머리 2
    notes = sum(1 for ws in wb.worksheets if ws.title.startswith("2.시험계획-")
                for c in ws["A"] if isinstance(c.value, str) and c.value.startswith("근거 기준"))
    assert body - notes - sum(1 for ws in wb.worksheets if ws.title.startswith("2.시험계획-")) == len(plan)   # 표 뒤 빈 줄 1


def test_hwpx_is_valid_and_has_parts(started):
    with zipfile.ZipFile(started / "품질시험계획서.hwpx") as z:
        text = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("Contents/section"))
    for part in ("1. 개요", "2. 품질시험 및 검사계획", "3. 품질시험 시설", "4. 품질관리자 배치계획", "(스캔첨부)"):
        assert part in text
    if VALIDATE is not None:
        run_hwpx_validate(started / "품질시험계획서.hwpx")


def test_format_and_out_dir(started, tmp_path):
    out = tmp_path / "only_xlsx"
    r = _run(["test-plan", "--project", str(started / "project.yaml"), "--format", "xlsx", "--out-dir", str(out),
              "--offline", "--json"], tmp_path / "h")
    assert r.returncode == 0, r.stderr
    assert sorted(p.name for p in out.iterdir()) == ["품질시험계획서.json", "품질시험계획서.xlsx"]
    assert json.loads(r.stdout)["hwpx"] is None
    r = _run(["test-plan", "--project", str(started / "project.yaml"), "--format", "pdf", "--out-dir", str(out)],
             tmp_path / "h")
    assert r.returncode == 2 and "--format" in r.stderr


def test_revision_mismatch_keeps_files(started, tmp_path):
    out = tmp_path / "rev"
    r = _run(["test-plan", "--project", str(started / "project.yaml"), "--revision", "0", "--date", "1999. 01. 01.",
              "--out-dir", str(out), "--offline"], tmp_path / "h")
    assert r.returncode == 2 and "시험계획서를 만들지 않았습니다" in r.stderr
    assert not out.exists() or not any(out.iterdir())


def test_missing_boq_path_is_input_error(started, tmp_path):
    proj = yaml.safe_load((started / "project.yaml").read_text(encoding="utf-8"))
    proj.pop("내역서")
    p = tmp_path / "project.yaml"
    p.write_text(yaml.safe_dump(proj, allow_unicode=True), encoding="utf-8")
    r = _run(["test-plan", "--project", str(p), "--offline"], tmp_path / "h")
    assert r.returncode == 2 and "내역서 경로가 없습니다" in r.stderr


def test_out_dir_inside_repo_is_refused(started, tmp_path):
    root = Path(__file__).resolve().parents[1]
    r = _run(["test-plan", "--project", str(started / "project.yaml"), "--out-dir", str(root / "src"), "--offline"],
             tmp_path / "h")
    assert r.returncode == 2 and "저장소" in r.stderr
    assert not (root / "src" / "품질시험계획서.hwpx").exists()


def test_quality_test_plan_site_does_not_overwrite_full_document(started, tmp_path):
    """품질시험계획 대상 현장: start 가 계획서 전체를 '품질시험계획서.hwpx' 로 쓴다 → 단독본은 '(단독)' 이름."""
    proj = yaml.safe_load((started / "project.yaml").read_text(encoding="utf-8"))
    proj["판정"]["계획종류"] = "품질시험계획"
    d = tmp_path / "qt"
    d.mkdir()
    (d / "project.yaml").write_text(yaml.safe_dump(proj, allow_unicode=True), encoding="utf-8")
    (d / "품질시험계획서.hwpx").write_bytes(b"full")
    r = _run(["test-plan", "--project", str(d / "project.yaml"), "--offline"], tmp_path / "h")
    assert r.returncode == 0, r.stderr
    assert (d / "품질시험계획서.hwpx").read_bytes() == b"full"
    assert (d / "품질시험계획서(단독).hwpx").exists() and (d / "품질시험계획서(단독).xlsx").exists()


def test_plain_input_start_then_test_plan(boq, tmp_path):  # noqa: F811
    """한 줄 모드 --input(에이전트 중계)으로 project.yaml 만 만든 뒤(계획서는 나중에), test-plan 이 질문 없이 만든다."""
    ans = tmp_path / "answers.txt"
    ans.write_text("\n".join(_lines(boq, "Y", "n")) + "\n", encoding="utf-8")
    out = tmp_path / "o"
    r = _run(["start", "--plain", "--input", str(ans), "--out-dir", str(out), "--folder", str(tmp_path)], tmp_path / "h")
    assert r.returncode == 0, r.stdout + r.stderr
    r = _run(["test-plan", "--project", str(out / "project.yaml"), "--offline", "--json"], tmp_path / "h", stdin="")
    assert r.returncode == 0, r.stdout + r.stderr
    s = json.loads(r.stdout)
    assert s["test_plan"]["rows"] > 0 and Path(s["hwpx"]).exists() and Path(s["xlsx"]).exists()


def test_owner_from_project_reaches_plan_and_test_plan_alike(started, boq, tmp_path):  # noqa: F811
    """project.yaml 발주처기준: LH → test-plan(--owner 없이)도 plan(--owner 없이)도 LH 로 계산하고 두 행이 같다(L14-D2 연결)."""
    proj = yaml.safe_load((started / "project.yaml").read_text(encoding="utf-8"))
    proj["발주처기준"] = "LH"
    d = tmp_path / "lh"
    d.mkdir()
    p = d / "project.yaml"
    p.write_text(yaml.safe_dump(proj, allow_unicode=True), encoding="utf-8")
    r = _run(["test-plan", "--project", str(p), "--offline", "--json"], tmp_path / "h")
    assert r.returncode == 0, r.stderr
    assert json.loads(r.stdout)["test_plan"]["owner"] == "LH"
    rev = proj["개정이력"][-1]
    r = _run(["plan", "--boq", str(boq), "--block", "나동", "--project", str(p), "--revision", str(rev["개정"]),
              "--date", str(rev["일자"]), "--offline", "--out", str(d / "plan" / "품질관리계획서.hwpx")], tmp_path / "h")
    assert r.returncode == 0, r.stderr
    lh_plan = json.loads((d / "plan" / "품질관리계획서.json").read_text(encoding="utf-8"))
    assert json.loads((d / "품질시험계획서.json").read_text(encoding="utf-8")) == lh_plan
    base = json.loads((started / "품질관리계획서.json").read_text(encoding="utf-8"))
    assert lh_plan != base                                           # LH 기준이 실제로 계산에 들어갔다
    assert any(str(x.get("basis", "")).startswith("LHCS") or "LHCS" in str(x.get("basis", "")) for x in lh_plan)


def test_default_output_is_human_result_screen(started, tp, tmp_path):
    """기본 출력 = 사람이 읽는 결과 화면(check 처럼). --json 요약과 같은 수치를 쓴다(L14-I2)."""
    out = tmp_path / "human"
    r = _run(["test-plan", "--project", str(started / "project.yaml"), "--out-dir", str(out), "--offline"], tmp_path / "h")
    assert r.returncode == 0, r.stderr
    lines = r.stdout.splitlines()
    assert lines[0].startswith("단번 품질시험계획서 — ")
    assert lines[1].startswith("만든 파일: 품질시험계획서.hwpx(한글) · 품질시험계획서.xlsx(엑셀) — 폴더 ")
    assert lines[2].startswith(f"시험 {tp['test_plan']['rows']}행 · {'·'.join(tp['test_plan']['disciplines'])} · 4부")
    assert any(ln.startswith("확인할 것") for ln in lines)
    assert any(ln.startswith("(스캔첨부) 빈칸 5곳: 시험실 배치평면도·경력증명서") for ln in lines)
    nxt = next(n for n, ln in enumerate(lines) if ln.startswith("다음 할 일"))
    assert lines[nxt] == "다음 할 일:" and lines[nxt - 2].startswith("(작성 필요)")        # start 가 남긴 (작성 필요) → 두 단계
    assert f'danburn test-plan --project "{started / "project.yaml"}"' in r.stdout      # 같은 명령·전체 경로
    assert "덮어쓰입니다" in r.stdout and "Claude Code" in lines[-1]
    assert not r.stdout.lstrip().startswith("{")
    assert 9 <= len(lines) <= 17                                          # 확인할 것은 3줄까지 보이고 나머지는 개수


def test_report_text_actions_and_single_step(tmp_path):
    """확인할 것 줄마다 할 일, (작성 필요)가 없으면 다음 할 일 한 줄(L14-I3)."""
    from danburn.cli import _testplan_report_text
    summary = {"warnings": ["발주처 기준 필요 — 시험계획 미작성: 합성재(별표2 밖, 내역 1행, 예: 합성재)",
                            "제조사(골재원) 수를 1곳으로 계산한 자재 1종 — 다르면 --makers 또는 project.yaml 제조사수",
                            "KS 인증 확인 전 — 합성을 KS 로 계산", "기타 합성 경고"],
               "test_plan": {"rows": 7, "disciplines": ["건축"], "scan_attachments": ["시험실 배치평면도", "재직증명서"]}}
    full = {k: "합성" for k in ("공사명", "공사위치", "공사금액", "공사기간", "발주자", "시공자", "건설사업관리자")}
    full.update(시험장비=[{"시험기구": "합성"}], 품질관리자=[{"성명": "합성"}], 조직=[{"성명": "합성"}])
    targets = {"hwpx": tmp_path / "품질시험계획서.hwpx", "xlsx": tmp_path / "품질시험계획서.xlsx"}
    text = _testplan_report_text(summary, full, targets, tmp_path, tmp_path / "project.yaml")
    lines = text.splitlines()
    assert lines[4].endswith("→ 발주처 시방서의 시험 기준을 받아 한글·엑셀 표에 직접 행 추가")
    assert lines[5].endswith("→ 제조사가 여럿이면 project.yaml 에 제조사수: {철근: 8} 처럼 적고 다시 실행")
    assert "--makers" not in lines[5]                                     # 명령어 말 대신 할 일 한 가지
    assert "KS_인증" in lines[6] and lines[7] == "  - 그 밖 1가지는 --json 의 warnings"
    assert not any(ln.startswith("(작성 필요)") for ln in lines)
    assert lines[-1].startswith("다음 할 일: 한글·엑셀로 열어")              # 한 줄
    assert "Claude Code" not in text

"""내역서 0행 원인 진단 — 합성 xlsx만 사용한다."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from danburn.boq import diagnose


def _book(path: Path, *, sheet="지급(건)", header_row=3, qty_name="수량",
          item_header=True, qty=None, formula=False):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in range(1, header_row):
        ws.cell(row, 1, "머리")
    if item_header:
        ws.cell(header_row, 1, "품명")
    ws.cell(header_row, 2, qty_name)
    ws.cell(header_row + 1, 1, "합성자재")
    ws.cell(header_row + 1, 2, "=1" if formula else qty)
    wb.save(path)
    return path


@pytest.mark.parametrize(
    ("filename", "kwargs", "expected"),
    [
        ("wrong-sheet.xlsx", {"sheet": "지급자재"}, "시트 이름이 달라요"),
        ("late-header.xlsx", {"header_row": 31}, "머리행을 30행 안에서 못 찾음"),
        ("no-name.xlsx", {"item_header": False}, "품명 칸 없음"),
        ("wrong-qty-name.xlsx", {"qty_name": "수량계"}, "수량 칸 이름이 '수량계'"),
        ("empty-qty.xlsx", {"qty": 0}, "수량이 전부 빈칸/0"),
        ("formula-qty.xlsx", {"formula": True}, "수식 결과가 저장 안 됨"),
    ],
)
def test_diagnose_synthetic_xlsx(tmp_path, filename, kwargs, expected):
    path = _book(tmp_path / filename, **kwargs)
    assert any(expected in reason for reason in diagnose(path))


def test_diagnose_legacy_xls(tmp_path):
    path = _book(tmp_path / "legacy.xlsx")
    legacy = path.with_suffix(".xls")
    path.rename(legacy)
    assert diagnose(legacy) == ["구형 .xls 파일이라 읽지 못함 → 엑셀에서 .xlsx로 다른 이름 저장"]

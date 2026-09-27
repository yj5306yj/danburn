"""danburn start 첫 실행 인터뷰(L7-I1): 경로 정리·답 파일·한 줄 모드·이어 하기·개인정보·저장소 안 경로 거부."""
import io
import json
import re
import struct
import subprocess
import sys
import zipfile
import zlib
from pathlib import Path

import pytest
import yaml

from danburn import start
from danburn.start import Interview, StartError, clean_path, tracked_in_repo
from conftest import VALIDATE, run_hwpx_validate

ROOT = Path(__file__).resolve().parents[1]
DANBURN = Path(sys.executable).parent / "danburn"


@pytest.fixture(scope="module")
def boq(tmp_path_factory):
    d = tmp_path_factory.mktemp("boq")
    p = d / "example_boq.xlsx"
    subprocess.run([sys.executable, str(ROOT / "scripts" / "make_example_boq.py"), "--out", str(p)],
                   check=True, capture_output=True)
    return p


def _png(path, color=(40, 90, 160)):
    """합성 단색 PNG(실제 문서 아님)."""
    w = h = 8
    raw = b"".join(b"\x00" + bytes(color) * w for _ in range(h))

    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    return path


def _answers(boq, **kw):
    a = {"내역서": str(boq), "블록": "나동", "발주자_구분": "민간", "발주자": "합성개발", "공사종류": "건축",
         "총공사비_억원": 850, "연면적": 42000, "지상층수": 22, "건설사업관리_대상": "아니오",
         "공사명": "합성 예시 공동주택", "시공자": "합성건설", "현장대리인": "합성 갑", "품질관리자": "합성 을"}
    a.update(kw)
    return a


def _run(args, env_home, stdin=None):
    env = {**__import__("os").environ, "DANBURN_HOME": str(env_home)}
    return subprocess.run([str(DANBURN), *args], input=stdin, capture_output=True, text=True, encoding="utf-8", env=env, timeout=300)


# ── 경로 정리 ─────────────────────────────────────────────────────────

@pytest.mark.parametrize("raw,want", [
    ("'/Users/a/내 파일.xlsx'", "/Users/a/내 파일.xlsx"),
    ('"/Users/a/b.xlsx"  ', "/Users/a/b.xlsx"),
    ("/Users/a/내\\ 파일\\(1\\).xlsx", "/Users/a/내 파일(1).xlsx"),      # 맥 터미널 끌어다 놓기
    ("file:///Users/a/%EB%82%B4.xlsx", "/Users/a/내.xlsx"),
    ('"C:\\Users\\a\\b.xlsx"', "C:\\Users\\a\\b.xlsx"),                  # 윈도 경로는 역슬래시 유지
])
def test_clean_path(raw, want):
    assert clean_path(raw) == want


# ── 답 파일(비대화형) ─────────────────────────────────────────────────

def test_answers_to_plan_end_to_end(boq, tmp_path):
    att = [_png(tmp_path / "합성_자격증_사본.png"), _png(tmp_path / "합성_경력증명서.png")]
    ans = tmp_path / "ans.yaml"
    ans.write_text(yaml.safe_dump(_answers(boq, 첨부=[str(p) for p in att]), allow_unicode=True), encoding="utf-8")
    out = tmp_path / "out"
    r = _run(["start", "--answers", str(ans), "--out-dir", str(out), "--offline"], tmp_path / "home")
    assert r.returncode == 0, r.stdout + r.stderr
    assert "품질관리계획" in r.stdout and "고급" in r.stdout
    for f in ("project.yaml", "품질관리계획서.hwpx", "품질관리계획서.json", "요약.json"):
        assert (out / f).exists(), f
    summary = json.loads((out / "요약.json").read_text(encoding="utf-8"))
    assert summary["start"]["첨부_개수"] == 2 and summary["start"]["판정"]["계획종류"] == "품질관리계획"
    # 개인정보: 첨부 이름·경로는 화면·요약에 없고 project.yaml 에만(절대 경로)
    for name in ("합성_자격증_사본", "합성_경력증명서"):
        assert name not in r.stdout + r.stderr and name not in (out / "요약.json").read_text(encoding="utf-8")
    proj = yaml.safe_load((out / "project.yaml").read_text(encoding="utf-8"))
    assert proj["첨부"]["파일"] == [str(p.resolve()) for p in att] and proj["첨부_개수"] == 2
    assert proj["승인절차_문장"].endswith(".") and "인·허가기관" in proj["승인절차_문장"]    # 민간
    assert proj["시험실"] == "50㎡ 이상" and proj["품질관리_대상등급"] == "고급"
    assert proj["판정"]["계획종류"] == "품질관리계획"
    if VALIDATE is not None:
        run_hwpx_validate(out / "품질관리계획서.hwpx")
    with zipfile.ZipFile(out / "품질관리계획서.hwpx") as z:
        text = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("Contents/section"))
    assert "시행령 제89조제1항제2호" in text and "인·허가기관의 장에게 제출" in text


def test_missing_required_answer_stops(boq, tmp_path):
    a = _answers(boq)
    del a["공사명"]
    ans = tmp_path / "a.yaml"
    ans.write_text(yaml.safe_dump(a, allow_unicode=True), encoding="utf-8")
    r = _run(["start", "--answers", str(ans), "--out-dir", str(tmp_path / "o"), "--no-plan"], tmp_path / "h")
    assert r.returncode == 2 and "공사명" in r.stdout


def test_quality_test_plan_sets_cover_title(boq, tmp_path):
    ans = tmp_path / "a.yaml"
    ans.write_text(yaml.safe_dump(_answers(boq, 총공사비_억원=30, 연면적=3000, 지상층수=5, 계약_품질관리계획="아니오",
                                           주용도="공동주택"), allow_unicode=True), encoding="utf-8")
    out = tmp_path / "o"
    r = _run(["start", "--answers", str(ans), "--out-dir", str(out), "--no-plan"], tmp_path / "h")
    assert r.returncode == 0, r.stdout
    proj = yaml.safe_load((out / "project.yaml").read_text(encoding="utf-8"))
    assert proj["판정"]["계획종류"] == "품질시험계획" and proj["문서명"] == "품질시험계획서"
    assert any("품질시험계획 대상" in n for n in proj["판정"]["확인필요"])


# ── 저장소 안 경로 거부 ───────────────────────────────────────────────

def test_tracked_in_repo(tmp_path):
    repo = tmp_path / "repo"
    repo.mkdir()
    subprocess.run(["git", "init", "-q", str(repo)], check=True)
    (repo / ".gitignore").write_text("danburn-out/\n")
    got = tracked_in_repo(repo / "산출")                                          # Windows: Git 은 C:/…, Path 는 C:\…
    assert got is not None and Path(got).resolve() == repo.resolve()
    assert tracked_in_repo(repo / "danburn-out" / "현장") is None                   # 무시되는 곳은 허용
    assert tracked_in_repo(tmp_path / "밖") is None


def test_out_dir_inside_repo_is_refused(boq, tmp_path):
    repo = tmp_path / "repo"                                  # 테스트 전용 저장소 — .git 없는 ZIP 배포본에서도 같은 뜻
    repo.mkdir()
    subprocess.run(["git", "init", "-q", str(repo)], check=True)
    ans = tmp_path / "a.yaml"
    ans.write_text(yaml.safe_dump(_answers(boq), allow_unicode=True), encoding="utf-8")
    r = _run(["start", "--answers", str(ans), "--out-dir", str(repo / "합성_산출_테스트"), "--no-plan"], tmp_path / "h")
    assert r.returncode == 2 and "저장소" in r.stdout
    assert not (repo / "합성_산출_테스트").exists()


def test_plan_file_given_as_boq_points_to_check(boq, tmp_path):
    """새로 만들기(start)에 계획서(hwp·hwpx·pdf)를 주면 기존 계획서 검사(danburn check)로 안내한다."""
    plan = tmp_path / "기존 계획서.hwpx"
    plan.write_bytes(b"PK")
    ans = tmp_path / "a.yaml"
    ans.write_text(yaml.safe_dump(_answers(boq, 내역서=str(plan)), allow_unicode=True), encoding="utf-8")
    r = _run(["start", "--answers", str(ans), "--no-plan", "--out-dir", str(tmp_path / "out")], tmp_path / "h")
    assert r.returncode == 2 and "danburn check" in r.stdout + r.stderr


# ── 한 줄 모드·대화형·이어 하기 ─────────────────────────────────────────

def _lines(boq, *extra, now=False, floors="22", logo=""):
    """질문 순서대로의 답(대화형·한 줄 모드 공통, L7-I3 흐름: 내역서 먼저 → 모호한 것만).

    내역서 → 동 → 공사명 → 발주자 → 발주청 여부 → 공사종류(읽은 값 확인 Y) → 총공사비·연면적·층수 → (조건부는 판정이
    갈릴 때만 — 22층·4.2만㎡면 없음) → 로고 → 회사명(시공사) → 사람 칸 묶음(기본 나중에)."""
    head = [str(boq), "3", "합성 예시 공동주택", "합성발주", "1", "Y", "850", "42000", floors, logo, "합성건설"]
    later = ["1", "", "", "합성 갑", "합성 을", "", ""] if now else [""]
    return [*head, *later, *extra]


def test_plain_mode_one_line_questions(boq, tmp_path):
    stdin = "\n".join(_lines(boq, "Y", "n")) + "\n"          # 판정 맞음, 계획서는 지금 안 만듦
    r = _run(["start", "--plain", "--out-dir", str(tmp_path / "o"), "--folder", str(tmp_path)], tmp_path / "h", stdin)
    assert r.returncode == 0, r.stdout + r.stderr
    qlines = [ln for ln in r.stdout.splitlines() if ln.startswith("Q")]
    assert qlines[0].startswith("Q1 내역서 |") and any(ln.startswith("Q2 발주자_구분 | ") and "1=예" in ln for ln in qlines)
    assert "JUDGE" in r.stdout and (tmp_path / "o" / "project.yaml").exists()


@pytest.mark.parametrize("enc", ["utf-8", "utf-8-sig", "cp949"])
def test_plain_input_file_replay_without_redirect(boq, tmp_path, enc):
    """--input: 답 목록 파일을 매번 처음부터 다시 돌리는 중계(셸 리다이렉션 없이 맥·Windows 같은 명령).
    빈 파일 → 첫 질문에서 대기(종료 2), 답을 쌓으면 다음 질문, 이어 하기 초안은 끼어들지 않는다."""
    home = tmp_path / "h"                                     # 같은 home 을 계속 써도 초안이 끼어들면 안 된다
    ans = tmp_path / "answers.txt"
    ans.write_bytes(b"")
    args = ["start", "--plain", "--input", str(ans), "--out-dir", str(tmp_path / "o"), "--folder", str(tmp_path)]
    r = _run(args, home)
    assert r.returncode == 2 and [ln for ln in r.stdout.splitlines() if ln.startswith("Q")][-1].startswith("Q1 내역서 |")
    assert "--input" in r.stdout and "초안에 저장" not in r.stdout          # 이 모드엔 초안이 없다 — 사실대로 안내
    ans.write_bytes((str(boq) + "\r\n").encode(enc))           # Windows 메모장식 줄끝·인코딩
    r = _run(args, home)
    assert r.returncode == 2 and "이어서" not in r.stdout      # 이어 하기 질문이 끼어들지 않는다
    assert [ln for ln in r.stdout.splitlines() if ln.startswith("Q")][-1].split(" | ")[0] != "Q1 내역서"
    ans.write_bytes(("\n".join(_lines(boq, "Y", "n")) + "\n").encode(enc))
    r = _run(args, home)
    assert r.returncode == 0, r.stdout + r.stderr
    assert (tmp_path / "o" / "project.yaml").exists() and not (home / "start-draft.yaml").exists()
    assert yaml.safe_load((tmp_path / "o" / "project.yaml").read_text(encoding="utf-8"))["공사명"] == "합성 예시 공동주택"


def test_interactive_help_unknown_and_resume(boq, tmp_path, monkeypatch):
    monkeypatch.setenv("DANBURN_HOME", str(tmp_path / "h"))
    draft = start.home_dir() / "start-draft.yaml"
    # 1차: 도움말(?)·잘못된 번호 → 다시 묻기, 발주자_구분(발주청 여부)에서 입력이 끊김
    out = io.StringIO()
    iv = Interview(mode="tty", folder=tmp_path, out=out,
                   inp=io.StringIO(f"{boq}\n3\n합성 예시 공동주택\n합성발주\n?\n9\n"))
    with pytest.raises(StartError):
        iv.run(draft=draft)
    text = out.getvalue()
    assert "이렇게 읽었습니다" in text and "도움말:" in text and "목록의 번호로 고르세요" in text
    saved = yaml.safe_load(draft.read_text(encoding="utf-8"))
    assert saved["answers"]["블록"] == "나동" and saved["next"] == "2"
    # 2차: 이어서(Y) → 발주청 여부부터, '-' = 모름
    rest = _lines(boq)[4:]
    rest[0] = "-"
    iv2 = Interview(mode="tty", folder=tmp_path, out=io.StringIO(), inp=io.StringIO("Y\n" + "\n".join(rest) + "\n"))
    ans = iv2.run(draft=draft)
    assert ans["발주자_구분"] == "모름" and ans["블록"] == "나동" and ans["공사명"] == "합성 예시 공동주택"
    assert ans["연면적"] == "42,000㎡" and ans["지상층수"] == 22


def test_attachments_listing_shows_count_only(tmp_path):
    _png(tmp_path / "합성_자격증.png")
    q = next(q for q in start.load_questions() if q["key"] == "첨부")
    out = io.StringIO()
    iv = Interview(mode="tty", folder=tmp_path, out=out, inp=io.StringIO(f"'{tmp_path / '합성_자격증.png'}'\n\n"))
    files = iv.ask(q, 1, 1)
    assert len(files) == 1 and "첨부 1건" in out.getvalue() and "합성_자격증" not in out.getvalue()


# ── plan_doc 하위 호환(판정 전 project.yaml) ───────────────────────────

def test_old_project_without_judged_keys_gets_defaults_and_note(tmp_path):
    from danburn.model import PlanRow
    from danburn.plan_doc import PRE_JUDGE_NOTE, build_plan
    p = yaml.safe_load((ROOT / "src" / "danburn" / "data" / "templates" / "project.example.yaml").read_text(encoding="utf-8"))
    for k in ("승인절차_문장", "품질관리_대상등급", "시험실"):
        p.pop(k, None)
    notes = []
    rows = [PlanRow("건축", "철근콘크리트공사", "레미콘 25-24-150", "슬럼프", 1000, "m3", "120㎥마다", "1,000/120", 9)]
    build_plan(rows, p, tmp_path / "x.hwpx", basis_version="합성", revision=0, date="2026. 01. 05.", notes=notes)
    assert notes == [PRE_JUDGE_NOTE]
    p["판정"] = {"승인절차_문장": "합성 절차 문장이다.", "품질관리_대상등급": "고급", "시험실": "50㎡ 이상"}   # 판정 블록만 있어도 펼침
    notes = []
    build_plan(rows, p, tmp_path / "y.hwpx", basis_version="합성", revision=0, date="2026. 01. 05.", notes=notes)
    assert notes == []


# ── L7-I2: 계획서 뒤 '현장 확인' 질문(요약 ask 목록, 엔진 L7-E7 형식을 가짜 요약으로) ──────────

def _ask(key, kind="rule_miss", **kw):
    return {"key": key, "label": f"합성 {key}", "kind": kind, "examples": [f"{key} 품명1", f"{key} 품명2"],
            "lines": 2, "units": ["m2"], "rule": f"rule_{key}", **kw}


def _ns(tmp_path, **kw):
    import argparse
    base = dict(answers=None, plain=True, yes=False, folder=str(tmp_path), out_dir=str(tmp_path / "o"),
                no_plan=False, offline=True)
    base.update(kw)
    return argparse.Namespace(**base)


class _FakePlan:
    """start.run_plan 대역: 부를 때마다 그때 project.yaml 을 기록하고 정해 둔 요약을 돌려준다."""

    def __init__(self, summaries):
        self.summaries, self.projects = list(summaries), []

    def __call__(self, args):
        proj = Path(args[args.index("--project") + 1])
        self.projects.append(yaml.safe_load(proj.read_text(encoding="utf-8")))
        return 0, json.loads(json.dumps(self.summaries.pop(0))), ""


def test_site_check_parsing():
    from danburn.start import _site_value, site_answers
    assert [_site_value(x) for x in ("1", "넣기", "예", "2", "빼기", "3", "나중에", "", "-")] == \
        ["예", "예", "예", "아니오", "아니오", None, None, None, None]
    with pytest.raises(StartError):
        _site_value("9")
    assert site_answers({"a": "넣기", "b": 2, "c": "나중에"}) == {"a": "예", "b": "아니오"}


def test_site_checks_asked_after_plan_then_rerun(boq, tmp_path, monkeypatch, capsys):
    items = [_ask("steel_fiber"), _ask("fiberboard"), _ask("ordinary_plywood", kind="install_only")]
    fake = _FakePlan([{"warnings": [], "ask": items}, {"warnings": [], "ask": [items[2]]}])
    monkeypatch.setattr(start, "run_plan", fake)
    monkeypatch.setenv("DANBURN_HOME", str(tmp_path / "h"))
    stdin = "\n".join(_lines(boq, "Y", "Y", "1", "2", "3")) + "\n"       # 판정 Y, 만들기 Y, 넣기·빼기·나중에
    monkeypatch.setattr(sys, "stdin", io.StringIO(stdin))
    assert start.main(_ns(tmp_path)) == 0
    out = capsys.readouterr().out
    assert "SITE 3" in out and "RERUN" in out
    q = [ln for ln in out.splitlines() if ln.startswith("Q확인 ")]
    assert len(q) == 3 and "1=넣기 2=빼기 3=나중에(모름)" in q[0] and "합성 steel_fiber" in q[0] and "steel_fiber 품명1" in q[0]
    assert "설치 내역" in q[2]                                                   # 설치 행에만 있음 문구
    assert len(fake.projects) == 2 and "현장_확인" not in fake.projects[0]
    assert fake.projects[1]["현장_확인"] == {"steel_fiber": "예", "fiberboard": "아니오"}   # 나중에는 저장 안 함
    saved = yaml.safe_load((tmp_path / "o" / "project.yaml").read_text(encoding="utf-8"))
    assert saved["현장_확인"] == {"steel_fiber": "예", "fiberboard": "아니오"}
    summary = json.loads((tmp_path / "o" / "요약.json").read_text(encoding="utf-8"))
    assert summary["start"]["현장_확인_남음"] == 1


def test_site_checks_all_later_means_no_rerun(boq, tmp_path, monkeypatch, capsys):
    fake = _FakePlan([{"warnings": [], "ask": [_ask("steel_fiber")]}])
    monkeypatch.setattr(start, "run_plan", fake)
    monkeypatch.setenv("DANBURN_HOME", str(tmp_path / "h"))
    monkeypatch.setattr(sys, "stdin", io.StringIO("\n".join(_lines(boq, "Y", "Y", "")) + "\n"))   # Enter = 나중에
    assert start.main(_ns(tmp_path)) == 0
    assert len(fake.projects) == 1 and "RERUN" not in capsys.readouterr().out


def test_many_site_checks_grouped_with_all_later(boq, tmp_path, monkeypatch, capsys):
    items = [_ask(f"k{i}") for i in range(6)] + [_ask(f"i{i}", kind="install_only") for i in range(4)]   # 10 > 8
    fake = _FakePlan([{"warnings": [], "ask": items}, {"warnings": [], "ask": []}])
    monkeypatch.setattr(start, "run_plan", fake)
    monkeypatch.setenv("DANBURN_HOME", str(tmp_path / "h"))
    # 첫 묶음(규칙 밖 6건) = 모두 나중에, 둘째 묶음(설치 4건) = 하나씩: 넣기·나중에·나중에·빼기
    stdin = "\n".join(_lines(boq, "Y", "Y", "2", "1", "1", "3", "3", "2")) + "\n"
    monkeypatch.setattr(sys, "stdin", io.StringIO(stdin))
    assert start.main(_ns(tmp_path)) == 0
    out = capsys.readouterr().out
    groups = [ln for ln in out.splitlines() if "하나씩 묻기" in ln]
    assert len(groups) == 2 and "규칙에 걸리지 않은 내역 쪽 6건" in groups[0]
    assert fake.projects[1]["현장_확인"] == {"i0": "예", "i3": "아니오"}


def test_answers_file_site_checks_applied_before_first_plan(boq, tmp_path, monkeypatch):
    ans = tmp_path / "a.yaml"
    ans.write_text(yaml.safe_dump(_answers(boq, 현장_확인={"steel_fiber": "넣기", "fiberboard": "나중에",
                                                        "ordinary_plywood": "아니오"}), allow_unicode=True),
                   encoding="utf-8")
    fake = _FakePlan([{"warnings": [], "ask": [_ask("fiberboard")]}])
    monkeypatch.setattr(start, "run_plan", fake)
    monkeypatch.setenv("DANBURN_HOME", str(tmp_path / "h"))
    assert start.main(_ns(tmp_path, answers=str(ans), plain=False)) == 0
    assert len(fake.projects) == 1                                              # 묻지 않고 한 번만 만든다
    assert fake.projects[0]["현장_확인"] == {"steel_fiber": "예", "ordinary_plywood": "아니오"}
    summary = json.loads((tmp_path / "o" / "요약.json").read_text(encoding="utf-8"))
    assert summary["start"]["현장_확인_남음"] == 1



# ── L7-I3: 내역서 먼저 → 모호한 것만, 조건부는 판정이 갈릴 때만, 로고는 별도 질문 ─────────────

def _plain(tmp_path, stdin, *args):
    return _run(["start", "--plain", "--out-dir", str(tmp_path / "o"), "--folder", str(tmp_path), "--no-plan", *args],
                tmp_path / "h", stdin)


def _qids(stdout):
    return [ln.split(" ", 1)[0][1:] for ln in stdout.splitlines() if ln.startswith("Q")]


def test_decided_case_skips_conditional_questions(boq, tmp_path):
    r = _plain(tmp_path, "\n".join(_lines(boq, "Y")) + "\n")        # 22층·4.2만㎡ → 영 89①2호로 확정
    assert r.returncode == 0, r.stdout + r.stderr
    ids = _qids(r.stdout)
    assert not any(i.startswith("5-") for i in ids)
    assert ids == ["1", "1-1", "6-1", "2-1", "2", "3", "4-1", "4-2", "4-3", "7", "6-4", "6"]
    assert "READ " in r.stdout and "ASKED 10 CONFIRMED 1" in r.stdout      # 첫 내역서 질문 빼고 10, 공사종류는 확인만


def test_split_case_asks_deciding_questions(boq, tmp_path):
    stdin = "\n".join(_lines(boq, floors="10")[:9] + ["2", "-", "1", "", "합성건설", "", "Y"]) + "\n"
    r = _plain(tmp_path, stdin)                                          # 10층 → 건설사업관리·계약·용도가 결과를 가른다
    assert r.returncode == 0, r.stdout + r.stderr
    ids = _qids(r.stdout)
    assert {"5-1", "5-2", "5-3"} <= set(ids)


def test_cover_values_read_and_confirmed(boq, tmp_path):
    import openpyxl
    wb = openpyxl.load_workbook(boq)
    ws = wb.create_sheet("표지", 0)
    ws["A3"] = "공 사 명 ： 합성 표지 공동주택 신축공사"                  # 전각 콜론도
    ws["A5"], ws["C5"] = "발 주 처", "합성토지주택공사"
    wb["원가계산서"]["A20"], wb["원가계산서"]["B20"] = "총 공 사 비", 85_000_000_000
    cover = tmp_path / "cover.xlsx"
    wb.save(cover)
    assert start.read_cover(str(cover)) == {"공사명": "합성 표지 공동주택 신축공사", "발주자": "합성토지주택공사",
                                            "총액": 85_000_000_000.0}
    stdin = "\n".join([str(cover), "3", "Y", "Y", "Y", "Y", "Y", "42000", "22", "", "", "", "Y"]) + "\n"
    r = _plain(tmp_path, stdin)
    assert r.returncode == 0, r.stdout + r.stderr
    assert "ASKED 6 CONFIRMED 5" in r.stdout and "관급자재가 빠졌을 수 있습니다" in r.stdout   # 동·연면적·층수·로고·회사명·묶음만 물음
    proj = yaml.safe_load((tmp_path / "o" / "project.yaml").read_text(encoding="utf-8"))
    assert (proj["공사명"], proj["발주자"], proj["발주자_구분"], proj["총공사비_억원"]) == \
        ("합성 표지 공동주택 신축공사", "합성토지주택공사", "발주청", 850.0)


def test_confirm_no_falls_back_with_read_value_as_default(boq, tmp_path):
    stdin = "\n".join(_lines(boq, "Y")[:5] + ["n", ""] + _lines(boq, "Y")[6:]) + "\n"   # 공사종류 n → Enter = 읽은 값
    r = _plain(tmp_path, stdin)
    assert r.returncode == 0, r.stdout
    assert "기본=건축+토목" in r.stdout
    proj = yaml.safe_load((tmp_path / "o" / "project.yaml").read_text(encoding="utf-8"))
    assert proj["공사종류"] == "건축+토목"


def _png_rgba(path, w=40, h=20):
    """합성 투명 PNG(가운데 사각형만 불투명 — 실제 로고 아님)."""
    rows = []
    for y in range(h):
        row = b"\x00"
        for x in range(w):
            inside = w // 4 <= x < 3 * w // 4 and h // 4 <= y < 3 * h // 4
            row += bytes((200, 40, 40, 255) if inside else (0, 0, 0, 0))
        rows.append(row)

    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 6, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(b"".join(rows))) + chunk(b"IEND", b""))
    return path


def test_logo_question_puts_logo_on_cover_and_every_header(boq, tmp_path):
    logo = _png_rgba(tmp_path / "합성_로고.png")
    r = _run(["start", "--plain", "--out-dir", str(tmp_path / "o"), "--folder", str(tmp_path), "--offline"],
             tmp_path / "h", "\n".join(_lines(boq, "Y", "Y", "", logo="1")) + "\n")   # 로고 = 폴더 후보 1번, 끝의 "" = 현장 확인 나중에
    assert r.returncode == 0, r.stdout + r.stderr
    assert any(ln.startswith("Q7 로고 |") and "1=합성_로고.png" in ln for ln in r.stdout.splitlines())
    proj = yaml.safe_load((tmp_path / "o" / "project.yaml").read_text(encoding="utf-8"))
    assert proj["로고"] == str(logo.resolve()) and proj["회사명"] == "합성건설"
    with zipfile.ZipFile(tmp_path / "o" / "품질관리계획서.hwpx") as z:
        secs = sorted((n for n in z.namelist() if n.startswith("Contents/section")),
                      key=lambda n: int("".join(c for c in n if c.isdigit())))
        xml = [z.read(n).decode("utf-8") for n in secs]
        bins = [n for n in z.namelist() if n.startswith("BinData/")]
    assert bins and "<hp:pic" in xml[0]                                         # 표지
    heads = [re.search(r"<hp:header .*?</hp:header>", x, re.S) for x in xml[1:]]
    assert all(h and "<hp:pic" in h.group(0) for h in heads)                    # 모든 쪽 머리


# ── L7-I4: 품명으로 자재를 알 수 없는 행(kind name_unknown, 엔진 L7-E8 형식) ─────────────

def _nu(name, guess="ductile_pipe", word="시멘트라이닝"):
    ch = {guess: f"합성 규칙 {guess}(으)로 넣기 — 규격의 '{word}' 가 맞을 때"} if guess else {}
    ch.update({"예": "시험 대상 자재 — 발주처 기준 필요로 표시", "아니오": "자재 아님(빼기)", "나중에": "답하지 않음"})
    return {"key": f"name:{name}", "label": name, "kind": "name_unknown", "examples": [name], "lines": 3,
            "units": ["M"], "rule": None, "spec_word": word, "specs": [word], "guess": {"key": guess}, "choices": ch}


def test_name_unknown_questions_and_answers(boq, tmp_path, monkeypatch, capsys):
    items = [_nu("합성 주철관 A"), _nu("합성 주철관 B"), _nu("합성 주철관 C"), _nu("합성 모름 D", guess=None), _nu("합성 E")]
    fake = _FakePlan([{"warnings": [], "ask": items}, {"warnings": [], "ask": []}])
    monkeypatch.setattr(start, "run_plan", fake)
    monkeypatch.setenv("DANBURN_HOME", str(tmp_path / "h"))
    # A=1(규칙으로 넣기) B=2(시험 대상) C=3(자재 아님) D=1(추정 없음 → 1=시험 대상) E=Enter(나중에)
    monkeypatch.setattr(sys, "stdin", io.StringIO("\n".join(_lines(boq, "Y", "Y", "1", "2", "3", "1", "")) + "\n"))
    assert start.main(_ns(tmp_path)) == 0
    out = capsys.readouterr().out
    q = [ln for ln in out.splitlines() if ln.startswith("Q확인 name:")]
    assert len(q) == 5
    assert "‘합성 주철관 A’은 품명으로 무슨 자재인지 알 수 없습니다(규격에 ‘시멘트라이닝’, 내역 3행)." in q[0]
    assert "1=합성 규칙 ductile_pipe(으)로 넣기" in q[0] and "2=시험 대상 자재(발주처 기준 필요) 3=자재 아님 4=나중에" in q[0]
    assert "1=시험 대상 자재(발주처 기준 필요) 2=자재 아님 3=나중에" in q[3] and "기본=나중에" in q[3]
    assert fake.projects[1]["현장_확인"] == {"name:합성 주철관 A": "ductile_pipe", "name:합성 주철관 B": "예",
                                           "name:합성 주철관 C": "아니오", "name:합성 모름 D": "예"}


def test_name_unknown_answers_file_accepts_rule_key(boq, tmp_path, monkeypatch):
    from danburn.start import site_answers
    assert site_answers({"name:합성 관": "steel_fiber", "name:합성 판": "빼기", "steel_fiber": "넣기",
                         "name:합성 모름": "나중에"}) == {"name:합성 관": "steel_fiber", "name:합성 판": "아니오",
                                                         "steel_fiber": "예"}
    with pytest.raises(StartError):
        site_answers({"steel_fiber": "ordinary_plywood"})                   # 종별 key 답은 name: 키에만
    ans = tmp_path / "a.yaml"
    ans.write_text(yaml.safe_dump(_answers(boq, 현장_확인={"name:합성 주철관 A": "ductile_pipe"}), allow_unicode=True),
                   encoding="utf-8")
    fake = _FakePlan([{"warnings": [], "ask": []}])
    monkeypatch.setattr(start, "run_plan", fake)
    monkeypatch.setenv("DANBURN_HOME", str(tmp_path / "h"))
    assert start.main(_ns(tmp_path, answers=str(ans), plain=False)) == 0
    assert fake.projects[0]["현장_확인"] == {"name:합성 주철관 A": "ductile_pipe"}


def test_many_with_name_unknown_group(boq, tmp_path, monkeypatch, capsys):
    items = [_ask(f"k{i}") for i in range(5)] + [_nu(f"합성 품명 {i}") for i in range(4)]     # 9 > 8
    fake = _FakePlan([{"warnings": [], "ask": items}, {"warnings": [], "ask": []}])
    monkeypatch.setattr(start, "run_plan", fake)
    monkeypatch.setenv("DANBURN_HOME", str(tmp_path / "h"))
    monkeypatch.setattr(sys, "stdin", io.StringIO("\n".join(_lines(boq, "Y", "Y", "2", "1", "3", "", "", "")) + "\n"))
    assert start.main(_ns(tmp_path)) == 0
    groups = [ln for ln in capsys.readouterr().out.splitlines() if "하나씩 묻기" in ln]
    assert any("품명으로 자재를 알 수 없는 내역 쪽 4건" in g for g in groups)
    assert fake.projects[1]["현장_확인"] == {"name:합성 품명 0": "아니오"}

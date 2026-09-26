"""설비 배관 토공을 토목으로 합산(L7-E3) — 합성 내역 행."""
from __future__ import annotations

from pathlib import Path

import pytest

from danburn.calc import aggregate, plan_rows
from danburn.model import BoqLine
from danburn.rules import load_rules

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def rules():
    return load_rules(ROOT / "src" / "danburn" / "data" / "rules")


def _line(disc, name, spec, qty, row, sheet=None, section="", supply="사급", block=""):
    sheet = sheet or {"건축": "내역(건)", "토목": "내역(토)", "기계": "내역(기)", "전기": "내역(전)"}[disc]
    return BoqLine(discipline=disc, sheet=sheet, row=row, name=name, spec=spec, unit="M3", qty=qty,
                   block=block, supply=supply, section=section)


def _by(mats, material):
    return {(m.discipline, m.spec): m for m in mats if m.material == material}


def test_mechanical_backfill_joins_civil_same_spec(rules):
    moved: list[dict] = []
    mats, _ = aggregate([
        _line("토목", "되메우기", "90%,양호", 500, 1, section="부대토목 > 토공사"),
        _line("기계", "되메우기", "90%,양호", 120, 2, section="기계설비 > 오배수공사"),
    ], rules, moved=moved)
    got = _by(mats, "backfill")
    assert set(got) == {("토목", "90%,양호")}
    m = got[("토목", "90%,양호")]
    assert m.qty == 620
    assert ("내역(기)", 2) in m.sources
    assert m.work == "토공사"                          # 설비 구분(오배수공사)은 공종 투표에 넣지 않음
    assert moved == [{"from": "기계", "material": "backfill", "spec": "90%,양호", "qty": 120, "unit": "m3", "lines": 1}]


def test_no_mechanical_rows_left_in_plan(rules):
    mats, _ = aggregate([_line("기계", "터파기", "토사", 80, 1), _line("기계", "되메우기", "90%,양호", 60, 2)], rules)
    rows = plan_rows(mats, rules)
    assert rows and {r.discipline for r in rows} == {"토목"}


def test_new_civil_row_when_civil_has_no_same_spec(rules):
    moved: list[dict] = []
    mats, _ = aggregate([
        _line("토목", "되메우기", "95%,양호", 500, 1),
        _line("기계", "되메우기", "90%,양호", 40, 2),
        _line("전기", "되메우기", "90%,양호", 10, 3),
    ], rules, moved=moved)
    got = _by(mats, "backfill")
    assert got[("토목", "95%,양호")].qty == 500
    assert got[("토목", "90%,양호")].qty == 50         # 기계+전기 → 토목 새 행
    assert {(x["from"], x["qty"]) for x in moved} == {("기계", 40), ("전기", 10)}


def test_architecture_earthwork_stays(rules):
    moved: list[dict] = []
    mats, _ = aggregate([_line("건축", "터파기", "토사", 300, 1), _line("토목", "터파기", "토사", 100, 2)], rules, moved=moved)
    got = _by(mats, "excavation_bearing")
    assert got[("건축", "토사")].qty == 300 and got[("토목", "토사")].qty == 100
    assert moved == []


def test_non_earthwork_mechanical_material_stays(rules):
    mats, _ = aggregate([BoqLine("기계", "내역(기)", 1, "레미콘", "25-24-150", "M3", 30.0, "", "사급")], rules)
    assert {m.discipline for m in mats} == {"기계"}


def test_block_follows_civil_all_blocks(rules):
    mats, _ = aggregate([
        _line("토목", "되메우기", "90%,양호", 200, 1),
        _line("기계", "되메우기", "90%,양호", 30, 2, block="B블록"),
    ], rules, block="B블록", all_blocks_for=frozenset({"토목"}))
    got = _by(mats, "backfill")
    assert got[("토목", "90%,양호")].qty == 230 and got[("토목", "90%,양호")].block == ""


def test_cli_summary_lists_moved_earthwork(tmp_path, capsys):
    import json

    import openpyxl

    from danburn.cli import main
    wb = openpyxl.Workbook()
    wb.active.title = "지급(토)"
    for title, rows in (("지급(토)", [["품명", "규격", "단위", "수량"], ["되메우기", "90%,양호", "M3", 500]]),
                        ("지급(기)", [["품명", "규격", "단위", "수량"], ["되메우기", "90%,양호", "M3", 70]])):
        ws = wb[title] if title in wb.sheetnames else wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    src = tmp_path / "e.xlsx"
    wb.save(src)
    assert main(["build", "--boq", str(src), "--out", str(tmp_path / "o.hwpx"), "--offline"]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert summary["earthwork_moved"] == [{"from": "기계", "material": "backfill", "spec": "90%,양호", "qty": 70,
                                           "unit": "m3", "lines": 1}]
    rows = json.loads((tmp_path / "o.json").read_text(encoding="utf-8"))
    assert {r["discipline"] for r in rows if r["material"] == "backfill"} == {"토목"}
    assert all(r["qty"] == 570 for r in rows if r["material"] == "backfill")

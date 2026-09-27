"""danburn 명령행: 도급내역서 → 8.11 JSON·HWPX.

예: danburn build --boq 지급자재내역서.xlsx --block 블록B --out out/8.11.hwpx
"""
from __future__ import annotations

import argparse
import io
import json
import re
import os
import sys
from pathlib import Path
from .paths import RULES_DIR

DEFAULT_RULES = RULES_DIR


class InputError(Exception):
    """사용자가 고칠 입력 문제(파일 없음·YAML 구문 등). 메시지 한 줄 = 어느 파일을 어떻게 고칠지. 종료 2."""


class OutputLocked(Exception):
    """기존 산출 파일을 바꾸지 못했다(Windows 에서 한글이 열어 둔 경우 등). 기존 파일은 그대로. 종료 2."""


# ── 입출력 인코딩(W01) ─────────────────────────────────────────────────

class _Utf8OrCp949Lines(io.TextIOBase):
    """파이프 stdin: 줄마다 UTF-8 로 읽고, 안 되면 cp949(한국어 Windows 실행기가 로캘 인코딩으로 보낸 답)로 읽는다."""

    def __init__(self, raw, keep=None):
        self._raw = raw
        self._keep = keep                                       # 바꾼 원래 stdin — 버려지면 buffer 까지 닫힌다

    def readable(self) -> bool:
        return True

    def isatty(self) -> bool:
        return False

    def readline(self, size: int = -1) -> str:
        b = self._raw.readline()
        if not b:
            return ""
        try:
            s = b.decode("utf-8")
        except UnicodeDecodeError:
            s = b.decode("cp949", errors="replace")
        s = s.lstrip("﻿")                                    # PowerShell 등이 붙이는 BOM
        return s[:-2] + "\n" if s.endswith("\r\n") else s

    def read(self, size: int = -1) -> str:
        return "".join(iter(self.readline, ""))


def _setup_stdio() -> None:
    """명령행 진입 때만: 파이프로 받는 stdout/stderr 는 UTF-8 로 쓴다(한국어 Windows 기본 cp949 는 '—' 등을 못 씀).
    TTY 콘솔은 파이썬이 이미 유니코드로 쓰고, PYTHONIOENCODING 을 명시했으면 그 설정을 따른다."""
    if os.environ.get("PYTHONIOENCODING"):
        return
    for name in ("stdout", "stderr"):
        s = getattr(sys, name)
        try:
            if not s.isatty() and (s.encoding or "").lower().replace("_", "-") not in ("utf-8", "utf8"):
                s.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError, OSError):            # 테스트 캡처·닫힌 스트림 등은 그대로
            pass
    try:
        if not sys.stdin.isatty() and hasattr(sys.stdin, "buffer"):
            sys.stdin = _Utf8OrCp949Lines(sys.stdin.buffer, keep=sys.stdin)
    except (AttributeError, ValueError, OSError):
        pass


# ── 입력 확인(W06) ────────────────────────────────────────────────────

def read_yaml(path: str | Path, what: str, flag: str) -> dict:
    """YAML 파일 → dict. 없거나 구문이 깨졌으면 어느 파일 몇째 줄을 고칠지 InputError 로."""
    import yaml
    p = Path(path)
    if not p.is_file():
        raise InputError(f"{what} 파일이 없습니다: {p} — {flag} 경로를 확인하세요(공백이 있으면 따옴표로 감싸기).")
    data = p.read_bytes()
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp949", errors="replace")            # 메모장 ANSI(한국어 Windows) 저장본
    try:
        doc = yaml.safe_load(text)
    except yaml.YAMLError as e:
        mark = getattr(e, "problem_mark", None)
        where = f" {mark.line + 1}째 줄" if mark is not None else ""
        why = getattr(e, "problem", None) or type(e).__name__
        raise InputError(f"{what} 파일의 YAML 구문이 깨졌습니다: {p}{where} ({why}) — 그 줄의 괄호·따옴표·콜론을 고친 뒤 "
                         "다시 실행하세요.") from None
    if doc is None:
        return {}
    if not isinstance(doc, dict):
        raise InputError(f"{what} 파일은 '키: 값' 목록이어야 합니다: {p} — 예제 파일 형식을 따라 고치세요.")
    return doc


def _check_boq_paths(paths) -> None:
    for s in paths or []:
        p = Path(s)
        if not p.exists():
            raise InputError(f"내역서 파일이 없습니다: {p} — --boq 경로를 확인하거나 내역서를 다시 고르세요"
                             "(공백이 있으면 따옴표로 감싸기).")
        if not p.is_file():
            raise InputError(f"내역서 경로가 파일이 아닙니다: {p} — 폴더가 아니라 .xlsx 파일을 고르세요.")
        if p.suffix.lower() == ".xls":
            raise InputError(f"옛 엑셀 형식(.xls)은 읽지 못합니다: {p.name} — 엑셀에서 'Excel 통합 문서(.xlsx)'로 저장해 다시 고르세요.")
        if p.suffix.lower() not in (".xlsx", ".xlsm"):
            raise InputError(f"내역서는 엑셀 .xlsx 파일이어야 합니다: {p.name} — 도급내역서 .xlsx 를 고르세요.")


def _read_boq(p):
    """read_boq + 사용자 말 오류(손상·암호·잠금 파일)."""
    from .boq import read_boq
    try:
        return read_boq(p)
    except Exception as e:
        raise InputError(f"내역서를 읽지 못했습니다: {Path(p).name} ({type(e).__name__}) — 엑셀에서 열리는지, 암호가 걸려 "
                         "있지 않은지 확인하고 다시 저장해 고르세요.") from e


def _project(a: argparse.Namespace) -> dict | None:
    """--project YAML(한 번만 읽는다)."""
    if not getattr(a, "project", None):
        return None
    if getattr(a, "_project_doc", None) is None:
        a._project_doc = read_yaml(a.project, "현장 정보(--project)", "--project")
    return a._project_doc


# ── 산출 파일 교체(W03) ────────────────────────────────────────────────

def _tmp_for(dst: Path) -> Path:
    return dst.with_name(f".{dst.stem}.tmp{os.getpid()}{dst.suffix}")


def _commit(pairs: list[tuple[Path, Path]]) -> None:
    """임시 파일들을 제자리로 한꺼번에 옮긴다. 기존 파일을 먼저 옆으로 치워 두고, 하나라도 실패하면 모두 되돌린다
    (Windows 에서 한글이 열어 둔 파일은 이름을 바꿀 수 없다 → OutputLocked, 기존 파일 그대로)."""
    moved: list[tuple[Path, Path]] = []
    placed: list[Path] = []
    try:
        for _, dst in pairs:
            if dst.exists():
                bak = dst.with_name(f".{dst.name}.{os.getpid()}.bak")
                os.replace(dst, bak)
                moved.append((dst, bak))
        for tmp, dst in pairs:
            os.replace(tmp, dst)
            placed.append(dst)
    except OSError as e:
        for dst in placed:
            dst.unlink(missing_ok=True)
        for dst, bak in reversed(moved):
            os.replace(bak, dst)
        busy = next((dst for _, dst in pairs if dst not in placed), pairs[0][1])
        raise OutputLocked(f"출력 파일을 바꾸지 못했습니다: {busy} ({type(e).__name__}) — 한글·엑셀 등에서 열려 있으면 "
                           "파일을 닫고 다시 실행하세요. 기존 파일은 그대로 두었습니다.") from None
    for _, bak in moved:
        try:
            bak.unlink()
        except OSError:
            pass


YES, NO = ("예", "네", "yes", "y", "true", "1", "o"), ("아니오", "아니요", "no", "n", "false", "0", "x")


def _confirmations(a: argparse.Namespace) -> dict[str, bool | str]:
    """현장 확인 답(L7-E7): project.yaml '현장_확인: {종별 key: 예|아니오}' 위에 --confirm 'key=예,key2=아니오' 를 덮는다."""
    out: dict[str, bool | str] = {}
    items: list[tuple[str, object]] = []
    if getattr(a, "project", None):
        items += list((_project(a).get("현장_확인") or {}).items())
    for part in filter(None, (x.strip() for x in (getattr(a, "confirm", "") or "").split(","))):
        k, _, v = part.partition("=")
        items.append((k.strip(), v.strip()))
    for k, v in items:
        t = str(v).strip().lower() if not isinstance(v, bool) else ("yes" if v else "no")
        if str(k).startswith("name:") and t not in YES + NO and re.fullmatch(r"[a-z][a-z0-9_]*", t):
            out[str(k)] = t                             # 품명 확인 행을 그 종별 규칙으로 넣기(L7-E8): 'name:<품명>=steel_fiber'
            continue
        if t not in YES + NO:
            raise ValueError(f"현장 확인 '{k}' 의 답 '{v}' 을 읽을 수 없습니다. 예 또는 아니오로 적으세요.")
        out[str(k)] = t in YES
    return out


def _ks_answer(a: argparse.Namespace) -> str | None:
    """project.yaml 'KS_인증: 예|아니오|모름'(danburn start 3-1 질문) → '예'·'아니오'·'모름'. --project 가 없으면 None.
    키가 없거나 비었으면 '모름'(확인 전). 읽을 수 없는 값은 ValueError."""
    if not getattr(a, "project", None):
        return None
    v = _project(a).get("KS_인증")
    if v is None or str(v).strip() in ("", "모름", "-"):
        return "모름"
    t = ("yes" if v else "no") if isinstance(v, bool) else str(v).strip().lower()
    if t in YES:
        return "예"
    if t in NO:
        return "아니오"
    raise ValueError(f"project.yaml 의 KS_인증 '{v}' 을 읽을 수 없습니다. 예·아니오·모름 중 하나로 적으세요.")


def _ks_warning(rows, rules) -> str | None:
    """KS 인증 확인 전(모름)인데 KS 여부로 계산이 갈리는 행(KS 종별 묶음 행)이 있으면 경고 한 줄."""
    labels = list(dict.fromkeys(rules[r.material].label for r in rows
                                if r.material in rules and rules[r.material].group_tests and rules[r.material].ks_mark))
    if not labels:
        return None
    shown = "·".join(labels[:3]) + (f" 외 {len(labels) - 3}종" if len(labels) > 3 else "")
    rebar = "rebar" in {r.material for r in rows}
    return (f"KS 인증 확인 전 — {shown}을(를) KS 인증 제품으로 보고 KS 칸 ◎(시험 면제"
            + (", 철근은 제조회사·규격별 1회" if rebar else "") + ")로 계산했습니다. "
            "비KS 가 섞였으면 project.yaml 의 KS_인증 을 아니오 로 바꿔(또는 --non-ks) 다시 만드세요 — "
            + ("철근은 50톤마다, 그 밖은 " if rebar else "") + "별표2 빈도·제조회사별 외부 시험 횟수로 바뀝니다.")


def _name_groups(items: list[dict]) -> list[dict]:
    """품명 확인 항목을 (설치 행 여부, 추정 종별 목록)이 같은 것끼리 한 질문으로 묶는다(L7-E9).
    묶음 key = 'name:<첫 품명>' — 이 key 로 답하면 묶음의 모든 품명에 같은 답."""
    groups: dict[tuple, dict] = {}
    for u in items:
        gk = (bool(u.get("install")), tuple(g["key"] for g in u["guesses"]))
        g = groups.get(gk)
        if g is None:
            g = groups[gk] = {"key": f"name:{u['name']}", "names": [], "members": [], "lines": 0, "units": [],
                              "spec_words": [], "guesses": u["guesses"], "install": gk[0]}
        g["names"].append(u["name"])
        g["members"].append(u)
        g["lines"] += u["lines"]
        g["units"] += [x for x in u["units"] if x not in g["units"]]
        if u["spec_word"] not in g["spec_words"]:
            g["spec_words"].append(u["spec_word"])
    for g in groups.values():
        n = len(g["names"])
        g["label"] = g["names"][0] + (f" 외 {n - 1}품명" if n > 1 else "")
    return list(groups.values())


def cmd_build(a: argparse.Namespace) -> int:
    from .calc import aggregate, plan_rows, rows_to_json
    from .hwpx_out import build_811
    from .rules import load_rules

    if getattr(a, "project", None):                       # 쓰기 전에 현장 정보·개정 일자를 검사(W03 — 실패해도 기존 파일 그대로)
        from .plan_doc import _revisions
        try:
            _revisions(_project(a), a.revision, a.date)
        except (ValueError, TypeError, KeyError) as e:
            print(f"계획서를 만들지 않았습니다: {e}", file=sys.stderr)
            return 2
    rules = load_rules(a.rules)
    owner = (a.owner or "").upper()
    rules = {k: r for k, r in rules.items() if not r.owner or r.owner.upper() == owner}   # 발주처 전용 규칙은 --owner 일 때만
    lines = []
    for p in a.boq:
        lines += _read_boq(p)
    if not lines:
        import openpyxl
        sheets = {str(p): openpyxl.load_workbook(p, read_only=True).sheetnames for p in a.boq}
        print("읽을 수 있는 행이 없습니다. v0은 지급자재 내역서 양식(시트 '지급(건)'·'지급(기)'·'지급(토)')만 읽습니다.\n"
              f"받은 파일의 시트: {json.dumps(sheets, ensure_ascii=False)}", file=sys.stderr)
        return 2
    blocks = sorted({ln.block for ln in lines if ln.block})
    if a.block is not None:
        hit = [b for b in blocks if a.block in b]
        if not hit and blocks:
            print(f"--block '{a.block}' 과 맞는 블록이 없습니다. 있는 블록: {blocks}", file=sys.stderr)
            return 2
        if len(hit) > 1:
            print(f"--block '{a.block}' 이 여러 블록에 걸립니다: {hit}. 더 구체적으로 적으세요.", file=sys.stderr)
            return 2
    used = [ln for ln in lines if a.block is None or not ln.block or a.block in ln.block]
    all_blocks_for = frozenset({"토목"}) if a.civil_scope == "공구" else frozenset()
    try:
        confirm = _confirmations(a)
        ks_answer = _ks_answer(a)
    except ValueError as e:
        print(str(e), file=sys.stderr)
        return 2
    from .calc import forced_matches, rule_for_key
    # 묶음 답 펼치기(L7-E9): 품명 확인 질문은 여러 품명을 한 key 로 묶어 묻는다 → 답을 묶음의 모든 품명 key 로 옮긴다
    if any(k.startswith("name:") for k in confirm):
        from .calc import match_rule as _mr
        from .index import coverage as _cov
        _used = [ln for ln in lines if a.block is None or not ln.block or a.block in ln.block]
        for u in (_nu := _cov(_used, matched=lambda ln: _mr(ln, rules) is not None)["name_unknown"]):
            u["key"] = f"name:{u['name']}"
        for g in _name_groups(_nu):
            if g["key"] in confirm:
                for u in g["members"]:
                    confirm.setdefault(u["key"], confirm[g["key"]])
    forced = forced_matches(lines, rules, {k for k, v in confirm.items() if v is True},   # '예' → 그 종별 규칙으로 넣는다
                            by_name={k[5:]: v for k, v in confirm.items() if k.startswith("name:") and isinstance(v, str)})
    mats, unread = aggregate(lines, rules, a.block, all_blocks_for, moved=(earthwork_moved := []), forced=forced)
    if not mats:
        print("자재·규격으로 묶인 행이 없습니다(블록 이름 또는 품명 매칭 확인).", file=sys.stderr)
        return 2
    found = {m.material for m in mats}
    warnings = [f"{r.material}({r.label}): 입력에서 0행 — 이 자재의 8.11 행이 없습니다(사급 내역·블록 선택 확인)"
                for r in rules.values() if r.material not in found and r.warn_if_missing]
    # hate(루프 3): 규칙이 없는 자재가 조용히 빠지지 않게 — 별표2 색인으로 식별해 알린다
    from .index import coverage
    used_lines = [ln for ln in lines if a.block is None or not ln.block or a.block in ln.block]
    covered_keys = frozenset(k for r in rules.values() for k in r.index_keys)
    from .calc import match_rule as _match_rule
    # 행 단위 판정(L7-E4): 규칙 파일이 있어도 그 행이 규칙에 안 걸리면 드러낸다(조용한 누락 방지)
    cov = coverage(used_lines, covered_keys=covered_keys,
                   matched=lambda ln: id(ln) in forced or _match_rule(ln, rules) is not None)
    uncovered = cov["uncovered"]
    produced_keys = {k for m in mats for k in rules[m.material].index_keys}
    from .index import load_index
    produced_keys |= {e.key for e in load_index() if e.rule in {m.material for m in mats}}   # 색인 rule 로 이어진 종별(레미콘·철근)
    for u in cov["unmatched_in_covered"]:
        u["produced"] = u["key"] in produced_keys
    # 경고·미작성 표에는 8.11 행이 하나도 안 생긴 종별만(이미 행이 있는 종별의 나머지 행은 대개 시공·부속 행 — 요약에만)
    rule_miss = [u for u in cov["unmatched_in_covered"] if not u["produced"]]
    labor_only = [u for u in cov.get("labor_only", []) if u["key"] not in produced_keys]   # 이미 산출된 종별은 빼기(L4-T3)
    # 현장 확인(L7-E7): 사용자에게 물을 목록(규칙이 있는 종별만 — 답이 '예'면 넣을 수 있다), '아니오'는 경고에서 빼고 따로 남긴다
    def _ask(u, kind):
        r = rule_for_key(u["key"], rules)
        return {"key": u["key"], "label": u["label"], "kind": kind, "examples": u["examples"][:3], "lines": u["lines"],
                "units": u.get("units", []), "rule": f"{r.material}({r.label})" if r else None}
    # rule_miss 는 8.11 행이 하나도 없는 종별만 묻는다(L7-E9) — 행이 있는 종별의 나머지 행은 요약 unmatched_in_covered 에만.
    # 다른 자재가 규격 말로 섞이던 행은 E8 에서 name_unknown 으로 따로 묻는다.
    ask = [_ask(u, "rule_miss") for u in cov["unmatched_in_covered"] if u["key"] not in confirm and not u["produced"]]
    ask += [_ask(u, "install_only") for u in labor_only if u["key"] not in confirm]
    ask = [x for x in ask if x["rule"]]
    no_keys = {k for k, v in confirm.items() if v is False}
    confirmed_excluded = [{"key": u["key"], "label": u["label"], "lines": u["lines"], "examples": u["examples"][:3]}
                          for u in (*uncovered, *cov["unmatched_in_covered"], *labor_only) if u["key"] in no_keys]
    uncovered = [u for u in uncovered if u["key"] not in no_keys]
    rule_miss = [u for u in rule_miss if u["key"] not in no_keys]
    labor_only = [u for u in labor_only if u["key"] not in no_keys]
    # 품명으로 자재를 알 수 없는 행(L7-E8): 규격 말로 다른 자재 이름 아래 묶지 않고 품명 그대로 묻는다. 답 key 는 'name:<품명>'
    name_unknown = cov.get("name_unknown", [])
    for u in name_unknown:
        u["key"] = f"name:{u['name']}"
    name_unknown_confirmed = [u for u in name_unknown if confirm.get(u["key"]) is True]
    confirmed_excluded += [{"key": u["key"], "label": u["name"], "lines": u["lines"], "examples": u["specs"][:3]}
                           for u in name_unknown if confirm.get(u["key"]) is False]
    # 설치 행뿐인 품명은 추정 종별이 모두 이미 8.11 에 있으면 묻지 않는다(설치 행에만 있음과 같은 원칙 L4-T3) — 요약 name_unknown 에는 남음
    name_unknown_open = [u for u in name_unknown if u["key"] not in confirm
                         and not (u.get("install") and all(g["key"] in produced_keys for g in u["guesses"]))]
    open_groups = _name_groups(name_unknown_open)                      # 같은 추정끼리 한 질문(L7-E9)
    for g in open_groups:
        for u in g["members"]:
            u["group"] = g["key"]
    ask += [{"key": g["key"], "label": g["label"], "kind": "name_unknown", "examples": g["names"][:3], "names": g["names"],
             "lines": g["lines"], "units": g["units"], "rule": None, "spec_word": g["spec_words"][0], "spec_words": g["spec_words"],
             "specs": [x for u in g["members"] for x in u["specs"]][:3], "guess": g["guesses"][0] if g["guesses"] else None,
             "guesses": g["guesses"],
             "choices": {**{x["key"]: f"{r.label}(으)로 넣기 — 규격에 그 자재가 맞을 때"
                            for x in g["guesses"] if (r := rule_for_key(x["key"], rules))},
                         "예": "시험 대상 자재 — 발주처 기준 필요로 표시", "아니오": "자재 아님(빼기)", "나중에": "답하지 않음"}}
            for g in open_groups]
    warnings += [f"품명으로 자재를 알 수 없음 — 확인 필요: {g['label']}(규격에 '{'·'.join(g['spec_words'])}', 내역 {g['lines']}행)"
                 for g in open_groups]
    warnings += [f"발주처 기준 필요 — 시험계획 미작성: {u['name']}(현장 확인: 시험 대상 자재, 규격에 '{u['spec_word']}', 내역 {u['lines']}행)"
                 for u in name_unknown_confirmed]                              # '예' → 발주처 기준 필요로 표시(L7-E8)
    warnings += [f"규칙 없음 — 시험계획 미작성: {u['label']}(별표2 p.{u['page']}, 내역 {u['lines']}행, 예: {', '.join(u['examples'][:2])})"
                 for u in uncovered]
    warnings += [f"규칙 밖 행 — 확인 필요: {u['label']}(규칙은 있으나 걸리지 않은 내역 {u['lines']}행, 예: {', '.join(u['examples'][:2])})"
                 for u in rule_miss]
    warnings += [f"설치 행에만 있음 — 자재 포함 여부 확인: {u['label']}(내역 {u['lines']}행, 예: {', '.join(u['examples'][:2])})"
                 for u in labor_only]
    from .extra import flag_extras
    from .calc import ambiguous, match_rule
    unmatched = [ln for ln in used_lines if id(ln) not in forced and match_rule(ln, rules) is None]
    extras = flag_extras(unmatched)                                   # 규칙에 잡힌 행은 ‘미작성’으로 경고하지 않는다(hate 루프 4)
    amb = [{"name": ln.name, "rules": ambiguous(ln, rules)} for ln in used_lines if ambiguous(ln, rules)]
    warnings += [f"규칙 겹침 — 확인 필요: '{x['name']}' → {', '.join(x['rules'])} (가장 앞 규칙으로 계산)" for x in amb[:20]]
    handled = {k for m in mats for k in rules[m.material].extra_keys}      # 실제로 행을 만든 발주처 규칙의 항목만 경고에서 뺀다
    extras["owner_standard_needed"] = [x for x in extras["owner_standard_needed"] if x["key"] not in handled]
    warnings += [f"발주처 기준 필요 — 시험계획 미작성: {x['label']}(별표2 밖, 내역 {x['lines']}행, 예: {', '.join(x['examples'][:2])})"
                 for x in extras["owner_standard_needed"]]
    from .calc import parse_sets
    member_inferred = None
    formwork_arg = a.formwork_sets
    if not formwork_arg and not a.no_infer:
        from .member import infer_members, to_formwork_arg
        member_inferred = infer_members(lines, a.block)
        formwork_arg = to_formwork_arg(member_inferred)
    non_ks = a.non_ks or ks_answer == "아니오"                 # --non-ks 가 project.yaml KS_인증 보다 우선
    rows = plan_rows(mats, rules, parse_sets(a.sets_per_lot), parse_sets(formwork_arg), a.include_optional,
                     parse_sets(a.makers), non_ks)
    ks_warn = _ks_warning(rows, rules) if ks_answer == "모름" and not a.non_ks else None
    if ks_warn:
        warnings.insert(0, ks_warn)                             # start '확인할 것' 맨 앞에
    from .calc import added_spec_rows, missing_common_specs
    missing_common = missing_common_specs(mats)
    if a.add_spec:
        rows += added_spec_rows([s.strip() for s in a.add_spec.split(",") if s.strip()], rules, parse_sets(a.makers))
    forced_src = {(ln.sheet, ln.row) for ln in lines if id(ln) in forced}
    for r in rows:                                          # 현장 확인으로 들어온 내역 행이 섞인 8.11 행
        if forced_src & set(r.sources):
            r.note = (r.note + "; " if r.note else "") + "현장 확인으로 포함"
    confirmed_included = sorted({(k, rule_for_key(k, rules).material) for k, v in confirm.items() if v is True and rule_for_key(k, rules)}
                                | {(k, rule_for_key(v, rules).material) for k, v in confirm.items()
                                   if isinstance(v, str) and rule_for_key(v, rules)})
    out = Path(a.out)
    versions = sorted({r.basis_version for r in rules.values()})
    json_path = out.with_suffix(".json")
    json_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_out, tmp_json = _tmp_for(out), _tmp_for(json_path)   # 둘 다 만든 뒤에만 교체(W03)
    from .hwpx_out import unlisted_items
    unlisted = unlisted_items(uncovered, extras["owner_standard_needed"], extras["site_measurements"], rule_miss=rule_miss)
    unlisted += [{"kind": "owner_standard", "label": u["name"], "basis": "현장 확인", "lines": u["lines"]}
                 for u in name_unknown_confirmed]
    unlisted += [{"kind": "품명 확인 필요", "label": g["label"], "basis": f"규격에 '{'·'.join(g['spec_words'])}'", "lines": g["lines"],
                  "action": "품명으로 자재를 확인한 뒤 시험 대상이면 시험계획 작성"} for g in open_groups]
    try:
        tmp_json.write_text(json.dumps(rows_to_json(rows), ensure_ascii=False, indent=1), encoding="utf-8")
        if getattr(a, "project", None):
            from .plan_doc import build_plan
            project = dict(_project(a))
            if a.logo:
                project["로고"] = a.logo          # CLI 가 YAML 보다 우선
            if a.company:
                project["회사명"] = a.company
            try:
                build_plan(rows, project, tmp_out, basis_version=", ".join(versions), revision=a.revision, date=a.date,
                           unlisted=unlisted, notice_footer=a.notice_footer, notes=warnings)
            except (ValueError, FileNotFoundError) as e:      # 개정 모순·로고 파일 없음·자리표시 누락 등 — 파일을 만들지 않음
                print(f"계획서를 만들지 않았습니다: {e}", file=sys.stderr)
                return 2
        else:
            build_811(rows, tmp_out, basis_version=", ".join(versions), logo=a.logo, company=a.company, unlisted=unlisted,
                      notice_footer=a.notice_footer,
                      generated_note=f"자동 산출 초안 — 품질관리자가 현장 조건을 확인한 뒤 확정한다. 규격 미판독 행 {len(unread)}건."
                                     + "".join(f" [누락 경고] {w}" for w in warnings
                                               if not w.startswith(("규칙 없음", "규칙 밖 행", "발주처 기준 필요", "품명으로 자재를 알 수 없음"))))   # 표로 싣는다
        os.chmod(tmp_out, 0o644)
        os.chmod(tmp_json, 0o644)
        _commit([(tmp_out, out), (tmp_json, json_path)])
    except OutputLocked as e:
        print(str(e), file=sys.stderr)
        return 2
    finally:
        tmp_out.unlink(missing_ok=True)
        tmp_json.unlink(missing_ok=True)
    from .calc import optional_tests
    if a.offline or "1" in (os.environ.get("DANBURN_OFFLINE"), os.environ.get("QCPLAN_OFFLINE")):   # 옛 이름 호환(L7-N1)
        basis_check = {"status": "unknown", "message": "--offline: 기준 최신 여부를 확인하지 않음"}
    else:
        from .basis import check_basis
        basis_check = check_basis(rules)
    # 행별 비고에 흩어진 확인 항목을 한곳에(새 에이전트 스킬 시험 L5 지적)
    from collections import defaultdict
    confirm: dict[str, list[str]] = defaultdict(list)
    for r in rows:
        for part in filter(None, (x.strip() for x in (r.note or "").split(";"))):
            confirm[part].append(r.item)
    needs_confirmation = [{"note": k, "rows": len(v), "items": sorted(set(v))[:10]} for k, v in confirm.items()]
    summary = {"needs_confirmation": needs_confirmation, "ambiguous_lines": amb,
               "labor_only_materials": [{k: u.get(k) for k in ("key", "label", "lines", "examples")} for u in labor_only],
               "owner_standard_needed": extras["owner_standard_needed"],
               "site_measurements_to_confirm": extras["site_measurements"],
               "uncovered_materials": [{k: u.get(k) for k in ("key", "label", "page", "lines", "examples")} for u in uncovered],
               "unmatched_in_covered": [{k: u.get(k) for k in ("key", "label", "page", "lines", "examples", "produced")}
                                        for u in cov["unmatched_in_covered"]],
               "basis_check": {k: basis_check.get(k) for k in ("status", "message", "checked_at")}, "civil_scope": a.civil_scope, "missing_common_specs": missing_common, "member_inferred": member_inferred, "formwork_sets_used": formwork_arg, "basis_version": versions, "warnings": warnings,
               "ks": {"KS_인증": ks_answer, "non_ks": non_ks, "warned": bool(ks_warn)}, "optional_tests_excluded": optional_tests(rules),
               "boq_lines_used": len(used), "spec_groups": len(mats), "plan_rows": len(rows), "work_missing": sum(1 for r in rows if not r.work), "earthwork_moved": earthwork_moved,
               "ask": ask, "confirmed_included": [{"key": k, "rule": m, "lines": sum(1 for ln in used if id(ln) in forced and forced[id(ln)][0].material == m
                                                                and (ln.name == k[5:] if k.startswith("name:") else f"name:{ln.name}" not in confirm))}
                                                   for k, m in confirmed_included],
               "confirmed_excluded": confirmed_excluded,
               "name_unknown": [{**{k: u[k] for k in ("key", "name", "spec_word", "specs", "lines", "units", "install", "guess", "guesses")},
                                 "shown": u in name_unknown_open, "group": u.get("group")} for u in name_unknown],
               "name_unknown_confirmed": [{"key": u["key"], "name": u["name"], "lines": u["lines"]} for u in name_unknown_confirmed], "unread_spec_lines": len(unread),
               "hwpx": str(out), "json": str(json_path)}
    print(json.dumps(summary, ensure_ascii=False))
    return 0


def cmd_check_basis(a: argparse.Namespace) -> int:
    """규칙 데이터의 기준 판이 현행 고시와 같은지 공개 법령 미러로 확인한다(사용자 파일은 보내지 않음)."""
    from .basis import check_basis
    from .rules import load_rules
    rules = load_rules(a.rules)
    if getattr(a, "offline", False) or "1" in (os.environ.get("DANBURN_OFFLINE"), os.environ.get("QCPLAN_OFFLINE")):
        def _no_net(url):
            raise OSError("offline")
        res = check_basis(rules, fetch=_no_net)            # 네트워크 없이: 규칙의 기준 판만 보이고 status unknown(종료 4)
        res["message"] = "--offline: 기준 최신 여부를 확인하지 않음 — 규칙 데이터의 판만 표시"
    else:
        res = check_basis(rules)
    print(json.dumps(res, ensure_ascii=False))
    return {"current": 0, "outdated": 3}.get(res["status"], 4)


_PLAN_STATUS = {"outdated": "개정됨", "unknown": "확인 필요", "current": "현행", "info": "참고"}


def _plan_report_text(path: str, doc, res: dict) -> str:
    """check 결과를 사람이 바로 읽는 요약으로: 개정됨 → 확인 필요 → 현행 순, 항목마다 계획서 위치와 할 일.
    판정하지 않은 참고(판 표시 없는 인용)는 한 줄로 묶는다."""
    fnd = [f for f in res.get("findings", []) if f.get("kind") != "ks"]
    ks = [f for f in res.get("findings", []) if f.get("kind") == "ks"]
    n = {s: sum(1 for f in fnd if f.get("status") == s) for s in ("outdated", "unknown", "current")}
    head = {"outdated": f"개정·폐지된 기준 {n['outdated']}건이 있습니다",
            "current": ("판정한 인용은 모두 현행입니다 — 단, 도구가 판정하지 못한 인용 "
                        f"{n['unknown']}건은 직접 확인하세요" if n["unknown"] else "찾은 기준 인용이 모두 현행입니다"),
            "unknown": "현행 여부를 판정하지 못했습니다"}.get(res.get("status"), "현행 여부를 판정하지 못했습니다")
    n_info = sum((f.get("count") or 1) for f in fnd if f.get("status") == "info")
    lines = [f"단번 기준 검사 — {Path(path).name} ({doc.format.upper()})",
             f"결과: {head}  (개정됨 {n['outdated']} · 확인 필요 {n['unknown']} · 현행 {n['current']}"
             + (f" · 연도·번호 표기가 없어 확인 못 한 곳 {n_info})" if n_info else ")"),
             f"기준표 확인일 {res.get('snapshot_checked_at') or '-'} · 검사일 {res.get('checked_at') or '-'}"]
    if res.get("message"):
        lines.append(res["message"])
    for status in ("outdated", "unknown", "current"):
        for f in (x for x in fnd if x.get("status") == status):
            cur = f" → 현행 {f['current']}" if f.get("current") and status != "current" else ""
            cnt = f" ({f['count']}곳)" if (f.get("count") or 1) > 1 else ""
            lines.append(f"\n[{_PLAN_STATUS[status]}] {f.get('cited') or f.get('norm')}{cur}{cnt}")
            if status in ("outdated", "unknown") and f.get("where"):
                lines.append(f"  계획서: “{f['where']}”")
            if status in ("outdated", "unknown") and f.get("advice"):
                lines.append(f"  할 일: {f['advice']}")
            elif status == "current" and f.get("advice"):             # 현행으로 본 이유(예: 시행일 전 날짜지만 현행 번호 인용)
                lines.append(f"  ({f['advice']})")
    info = [f for f in fnd if f.get("status") == "info"]
    if info:
        names = ", ".join(f"{f.get('cited') or f.get('norm')}" + (f"({f['count']}곳)" if (f.get("count") or 1) > 1 else "")
                          for f in info)
        lines.append(f"\n참고: 연도·번호 표기 없이 이름만 적혀 확인 못 한 곳 — {names}")
    if ks:
        lines.append(f"참고: KS 번호 {len(ks)}개는 개정·폐지를 판정하지 않았습니다 — e-나라표준인증(standard.go.kr)에서 확인하세요.")
    for w in getattr(doc, "warnings", []) or []:
        lines.append(f"주의(파일 읽기): {w}")
    lines.append("\n이 결과는 알림입니다. 계획서를 고치기 전에 공식 원문(법제처·국가건설기준센터)으로 확인하세요.")
    return "\n".join(lines)


def cmd_check(a: argparse.Namespace) -> int:
    """완성된 품질관리계획서(HWP·HWPX·PDF)가 인용한 기준의 판이 현행인지 본다. 계획서 내용은 네트워크로 보내지 않는다."""
    from .plancheck import check_plan, load_snapshot
    from .readdoc import UnreadableDocument, read_document
    try:
        doc = read_document(a.file)
    except (UnreadableDocument, FileNotFoundError) as exc:
        print(json.dumps({"status": "unreadable", "message": str(exc)}, ensure_ascii=False) if a.json
              else f"계획서를 읽지 못했습니다: {exc}")
        return 2
    fetch = None
    if a.offline or "1" in (os.environ.get("DANBURN_OFFLINE"), os.environ.get("QCPLAN_OFFLINE")):
        def fetch(url):                                       # 업무지침 현행은 동봉 기준표로 판정
            raise OSError("offline")
    res = check_plan(doc.text, snapshot=load_snapshot(), fetch=fetch)
    res["file"] = {"name": Path(a.file).name, "format": doc.format, "warnings": list(doc.warnings)}
    print(json.dumps(res, ensure_ascii=False) if a.json else _plan_report_text(a.file, doc, res))
    return {"current": 0, "outdated": 3}.get(res.get("status"), 4)


def cmd_inspect(a: argparse.Namespace) -> int:
    """입력 확인: 시트, 읽힌 행 수, 블록 목록. 스킬 2·3단계용."""
    import openpyxl
    from collections import Counter
    for p in a.boq:
        lines = _read_boq(p)
        sheets = openpyxl.load_workbook(p, read_only=True).sheetnames
        print(json.dumps({"file": str(p), "sheets": sheets, "readable_lines": len(lines),
                          "by_sheet": dict(Counter(ln.sheet for ln in lines)),
                          "blocks": dict(Counter(ln.block or "(공통)" for ln in lines)),
                          "supported": bool(lines)}, ensure_ascii=False))
    return 0


def _add_calc_args(b: argparse.ArgumentParser) -> None:
    b.add_argument("--boq", nargs="+", required=True, help="도급내역서 xlsx (지급자재 내역서 형식)")
    b.add_argument("--block", default=None, help="블록·공구 열 이름 일부 (예: 블록B). 없으면 전체")
    b.add_argument("--rules", default=str(DEFAULT_RULES))
    b.add_argument("--sets-per-lot", default="", help="레미콘 압축강도 로트당 조 수를 통째로 지정. 예: '25-24-150=7'. 없으면 규칙 기본값(28일3+7일1=4조)")
    b.add_argument("--owner", default="", help="발주처 기준 규칙을 켠다(예: LH → LHCS 10 40 00 부록 「품질시험 및 검사기준」). 없으면 별표2만")
    b.add_argument("--logo", default=None, help="내 회사 로고 이미지(PNG/JPEG, 로컬 파일). 없으면 로고 칸을 비워 둔다")
    b.add_argument("--company", default="", help="내 회사명(로고 칸 아래 표시). 없으면 비워 둔다")
    b.add_argument("--offline", action="store_true", help="기준 최신 여부 확인(공개 법령 미러 조회)을 건너뛴다")
    b.add_argument("--civil-scope", choices=["블록", "공구"], default="블록",
                   help="토목 물량 범위. 원칙은 블록별, 한 업체가 두 공구 토목을 함께 맡으면 '공구'(모든 블록 합계)")
    b.add_argument("--add-spec", default="", help="도급내역서에 빠진 흔한 규격을 수량 없이 넣는다. 예: '건축:rebar:SD400 D10,건축:rebar:SD400 D13'")
    b.add_argument("--makers", default="", help="철근 제조회사 수(규격별). 예: 'SD400 D13=2,*=1'. 없으면 1곳+확인 필요")
    b.add_argument("--non-ks", action="store_true", help="KS 인증품이 아니면(철근 50톤마다 등 외부 시험). plan 의 project.yaml KS_인증 보다 우선")
    b.add_argument("--no-infer", action="store_true", help="부위 자동 추정을 끈다(거푸집 해체용 조를 더하지 않음)")
    b.add_argument("--include-optional", action="store_true", help="조건부·실무 생략 시험(휨강도, 온도·배합설계·현장배합수정)도 넣는다")
    b.add_argument("--formwork-sets", default="", help="거푸집 해체용 조를 규격별로 더한다. 부위로: 기둥·기초는 수직+예비, 슬래브·보까지면 수직+수평+예비. 예: '25-24-150=수직+수평+예비,25-24-80=수직+예비'")
    b.add_argument("--out", required=True, help="출력 HWPX 경로 (같은 이름 .json 도 생성)")
    b.add_argument("--confirm", default="", help="현장 확인 답(요약 ask 목록의 key). 예: 'steel_fiber=예,fiberboard=아니오' — 예: 그 종별 규칙으로 8.11 에 넣음, 아니오: 경고에서 뺌. project.yaml '현장_확인' 보다 우선")
    b.add_argument("--notice-footer", action="store_true", help="쪽 아래에 고지 줄 「본 제품은 한글과컴퓨터의 HWP 문서 파일(.hwp) 공개 문서를 참고하여 개발하였습니다.」를 넣는다(기본 끔 — 제출 문서용, 고지는 랜딩·README·NOTICE 에 둔다)")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="danburn")
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build", help="도급내역서에서 8.11 초안(HWPX+JSON)을 만든다")
    _add_calc_args(b)
    b.set_defaults(func=cmd_build)
    pl = sub.add_parser("plan", help="도급내역서+현장 정보로 품질관리계획서 전체(표지·목차·개정이력·1~10장, 8.11 계산 표)를 만든다")
    _add_calc_args(pl)
    pl.add_argument("--project", required=True, help="현장 정보 YAML (예: src/danburn/data/templates/project.example.yaml)")
    pl.add_argument("--revision", default="0", help="개정 번호 (예: 0, 3)")
    pl.add_argument("--date", required=True, help="개정 일자 (예: 2026. 09. 26.)")
    pl.set_defaults(func=cmd_build)
    c = sub.add_parser("check-basis", help="기준 판(업무지침 고시 번호)이 현행인지 확인한다. 종료 0=현행, 3=개정됨, 4=확인 못 함")
    c.add_argument("--rules", default=str(DEFAULT_RULES))
    c.add_argument("--offline", action="store_true", help="공개 법령 미러를 조회하지 않고 규칙 데이터의 기준 판만 보인다(종료 4)")
    c.set_defaults(func=cmd_check_basis)
    cp = sub.add_parser("check", help="기존 계획서 검사: 이미 있는 품질관리계획서(HWP·HWPX·PDF)가 인용한 기준이 현행인지 본다. 종료 0=현행, 3=개정됨, 4=판정 못 함, 2=파일을 못 읽음")
    cp.add_argument("file", help="검사할 계획서 파일(.hwp .hwpx .pdf)")
    cp.add_argument("--json", action="store_true", help="사람용 요약 대신 JSON 으로 출력(에이전트·스크립트용)")
    cp.add_argument("--offline", action="store_true", help="공개 법령 미러를 조회하지 않고 동봉 기준표로만 판정")
    cp.set_defaults(func=cmd_check)
    i = sub.add_parser("inspect", help="도급내역서를 읽을 수 있는지, 어떤 블록이 있는지 확인한다")
    i.add_argument("--boq", nargs="+", required=True)
    i.set_defaults(func=cmd_inspect)
    st = sub.add_parser("start", help="새로 만들기: 도급내역서와 몇 가지 질문으로 품질관리계획서 초안을 만든다(처음이면 여기부터)")
    st.add_argument("--answers", default=None, help="질문 답을 담은 yaml(비대화형). 키는 src/danburn/data/interview.yaml 의 key")
    st.add_argument("--plain", action="store_true", help="한 줄 질문·번호 응답(에이전트·스크린리더용)")
    st.add_argument("--input", default=None, help="한 줄 모드 답을 표준입력 대신 이 파일(한 줄에 답 하나)에서 읽는다. "
                    "에이전트 중계용 — 셸 리다이렉션(<) 없이 맥·Windows 같은 명령. 이어 하기 초안은 쓰지 않는다")
    st.add_argument("--yes", action="store_true", help="--answers 에 없는 답은 기본값(없으면 모름)으로")
    st.add_argument("--folder", default=".", help="내역서·로고 후보를 찾을 폴더(기본: 지금 폴더)")
    st.add_argument("--out-dir", default=None, help="산출 폴더(기본: ~/Documents/danburn/<공사명>/, 저장소 밖)")
    st.add_argument("--no-plan", action="store_true", help="project.yaml 만 만들고 계획서는 만들지 않는다")
    st.add_argument("--offline", action="store_true", help="기준 최신 여부 확인(공개 법령 미러 조회)을 건너뛴다")
    st.set_defaults(func=cmd_start)
    if argv is None:
        _setup_stdio()                                   # 명령행 진입만(같은 프로세스 호출·테스트 캡처는 건드리지 않음)
    if argv is None and len(sys.argv) == 1 or argv == []:
        return _no_args(ap)
    a = ap.parse_args(argv)
    try:
        _check_boq_paths(getattr(a, "boq", None))
        return a.func(a)
    except InputError as e:
        print(f"입력 오류: {e}", file=sys.stderr)
        return 2
    except Exception as e:                               # 예상 못 한 오류: 한 줄로(상세는 DANBURN_DEBUG=1)
        if os.environ.get("DANBURN_DEBUG") == "1":
            raise
        print(f"예상하지 못한 오류로 멈췄습니다: {type(e).__name__}: {e} — 자세한 내용은 DANBURN_DEBUG=1 로 다시 실행해 "
              "보고해 주세요.", file=sys.stderr)
        return 1


def cmd_start(a: argparse.Namespace) -> int:
    from .start import main as start_main
    return start_main(a)


def _no_args(ap: argparse.ArgumentParser) -> int:
    """인자 없이 실행: 터미널이면 두 갈래(새로 만들기 start / 기존 계획서 검사 check) 중 고르게 하고, 아니면 도움말."""
    if not sys.stdin.isatty():
        ap.print_help()
        return 0
    print("danburn — 품질관리계획서")
    print("  1) 새로 만들기     도급내역서(xlsx)로 초안부터       danburn start")
    print("  2) 기존 계획서 검사  가진 계획서(hwp·hwpx·pdf)의 기준이 현행인지   danburn check <파일>")
    choice = input("번호를 고르세요 [1/2, Enter=1, q=그만] > ").strip().lower()
    if choice in ("q", "n", "no", "아니오"):
        ap.print_help()
        return 0
    if choice == "2":
        from .start import clean_path
        path = clean_path(input("검사할 계획서 파일을 끌어다 놓거나 경로를 적으세요 > "))
        return main(["check", path]) if path else 2
    return main(["start"])


if __name__ == "__main__":
    sys.exit(main())

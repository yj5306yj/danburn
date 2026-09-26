"""도급내역서(xlsx) → BoqLine.

읽는 시트:
- 지급자재 `지급(건)`/`지급(기)`/`지급(토)` → supply="지급"
- 사급 내역 `내역(건)`·`내(건)` 등(뒤에 블록 꼬리가 붙어도 됨) → supply="사급"
원가·지구예산·품질관리비 등 그 밖의 시트는 조용히 건너뛴다.

머리 구조(병합 셀)는 파일마다 조금씩 다르므로 위치를 고정하지 않고 찾는다.
- "수량" 머리 칸마다 그 위 머리 행들에서 블록명, 변경 묶음(변경 전/후/증감), 금액 묶음(설계·도급·하도급)을 읽는다.
- 금액 묶음이 여럿이면 "도급 금액"을 쓴다. 변경 묶음이 여럿이면 "변경 후"를 쓴다("변경 후"가 비었으면 "변경 전").
- 합계 열, 목차·구분·소계 행, 수량이 0이거나 빈 행은 버린다.
- 구분 행("1-1 0102. 철근콘크리트공사" 같은 번호 제목)의 계층을 따라 행마다 section 경로를 채운다.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from .model import BoqLine

_DISCIPLINE = {"건": "건축", "기": "기계", "토": "토목"}
_SHEETS = (
    (re.compile(r"^\s*지급\s*\(\s*([건기토])\s*\)"), "지급"),
    (re.compile(r"^\s*내(?:역)?\s*\(\s*([건기토])\s*\)"), "사급"),
)
_MARKER = re.compile(r"(목차|구분|소계|합계|총계)")
_TOTAL_NAMES = {"계", "합계", "소계", "총계"}
_TOTAL_WORDS = ("합계", "총계", "소계")
_HEADING = re.compile(r"^\s*(\d+(?:-\d+)*)(?:\s+(\d+))?\s*\.\s*(.*?)\s*$")
_CONCRETE = re.compile(r"^\s*(\d{2})\s*-\s*(\d{2})\s*-\s*(\d{1,3})(?!\d)")
_REBAR_NAME = re.compile(r"(봉강|철근)")
_REBAR_GRADE = re.compile(r"SD\s*(\d{3})", re.I)
_REBAR_DIA = re.compile(r"(?:^|[^A-Z])(?:H?D|H)\s*-?\s*(\d{2})(?!\d)", re.I)
_UNIT_ALIASES = {
    "m3": "m3", "㎥": "m3", "m³": "m3",
    "ton": "ton", "톤": "ton", "t": "ton",
    "m2": "m2", "㎡": "m2", "m²": "m2",
}
SECTION_SEP = " > "


def _norm(value) -> str:
    """머리 글자 비교용: 공백 제거."""
    return re.sub(r"\s+", "", str(value)) if value is not None else ""


def normalize_unit(unit: str) -> str:
    raw = (unit or "").strip()
    return _UNIT_ALIASES.get(raw.lower(), raw.lower())


def normalize_concrete_spec(spec: str) -> str | None:
    """레미콘 규격(굵은골재-강도-슬럼프)을 슬럼프 mm 표기로.

    슬럼프가 1~2자리이고 30 미만이면 cm로 보고 ×10 한다(슬럼프 30cm 이상은 없으므로
    "80" 같은 두 자리 값은 이미 mm).
    """
    m = _CONCRETE.match(spec or "")
    if not m:
        return None
    size, strength, slump = m.groups()
    if len(slump) <= 2 and int(slump) < 30:
        slump = str(int(slump) * 10)
    return f"{size}-{strength}-{slump}"


def normalize_rebar_spec(name: str, spec: str) -> str | None:
    """철근 재료 행이면 "SD500 D13" 꼴로. 강종은 품명·규격 어디든, 호칭지름은 규격의 H-13/D-13/D13에서.

    가공·조립·시공도 같은 노무 행은 강종이 없으므로 None.
    """
    name, spec = name or "", spec or ""
    if not _REBAR_NAME.search(name):
        return None
    grade = _REBAR_GRADE.search(name) or _REBAR_GRADE.search(spec)
    dia = _REBAR_DIA.search(spec)
    if not grade or not dia:
        return None
    return f"SD{grade.group(1)} D{int(dia.group(1))}"


def _qty(value) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


class _Grid:
    """병합 셀을 풀어 (행, 열) → 값을 돌려준다."""

    def __init__(self, ws):
        self.ws = ws
        self._owner: dict[tuple[int, int], tuple[int, int]] = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self._owner[(r, c)] = (rng.min_row, rng.min_col)

    def value(self, row: int, col: int):
        r, c = self._owner.get((row, col), (row, col))
        return self.ws.cell(r, c).value

    def bottom(self, row: int, col: int) -> int:
        """병합 범위의 마지막 행."""
        for rng in self.ws.merged_cells.ranges:
            if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
                return rng.max_row
        return row


def _find_header(grid: _Grid, max_scan: int = 30):
    """'수량' 칸이 있는 머리 행과 품명·규격·단위 열을 찾는다. 없으면 None."""
    ws = grid.ws
    for r in range(1, min(ws.max_row, max_scan) + 1):
        qty_cols = [c for c in range(1, ws.max_column + 1)
                    if _norm(ws.cell(r, c).value) == "수량"]
        if not qty_cols:
            continue
        cols = {}
        for rr in range(1, r + 1):
            for c in range(1, ws.max_column + 1):
                key = _norm(ws.cell(rr, c).value)
                if key in ("품명", "규격", "단위") and key not in cols:
                    cols[key] = c
        if "품명" not in cols:
            continue
        end = max(grid.bottom(r, c) for c in qty_cols)
        return r, end, qty_cols, cols
    return None


def _is_total(label) -> bool:
    """합계 열·행 머리: '합 계', '⑥ 합 계 [④+⑤]', '계' 등."""
    text = _norm(label)
    return text in _TOTAL_NAMES or any(w in text for w in _TOTAL_WORDS)


def _is_change_label(label: str) -> bool:
    return any(k in label for k in ("변경전", "변경후", "증감"))


def _is_marker_row(ws, row: int) -> bool:
    first = ws.cell(row, 1).value
    return first is not None and bool(_MARKER.search(_norm(first)))


def _has_qty(ws, cols: list[int], first_row: int) -> bool:
    return any(_qty(ws.cell(r, c).value)
               for r in range(first_row, ws.max_row + 1) if not _is_marker_row(ws, r)
               for c in cols)


def _pick(groups, word: str) -> str:
    hits = [g for g in groups if word in g]
    if len(hits) != 1:
        raise ValueError(f"수량 묶음을 고를 수 없음: {sorted(groups)}")
    return hits[0]


def _qty_columns(grid: _Grid, qty_row: int, qty_cols: list[int], data_row: int) -> list[tuple[int, str]]:
    """수량 열마다 (열, 블록명). 합계 열은 뺀다.

    금액 묶음(설계·도급·하도급 …)이 여럿이면 '도급 금액'(없으면 가장 왼쪽)만 쓴다.
    변경 묶음(변경 전/후/증감)이 있으면 '변경 후'만 쓴다. 다만 '변경 후' 열이 모두 비어 있으면
    (변경이 아직 입력되지 않은 양식) '변경 전'을 쓴다. 증감은 쓰지 않는다.
    """
    found = []   # (열, 블록, 변경 묶음, 금액 묶음)
    for c in qty_cols:
        block, group, money = None, None, None
        for r in range(qty_row - 1, 0, -1):
            raw = grid.value(r, c)
            label = _norm(raw)
            if not label:
                continue
            if _is_change_label(label):
                group = group or label
            elif "금액" in label:
                money = money or label
            elif block is None and "단가" not in label:
                block = str(raw).strip()
        found.append((c, block or "", group, money))

    monies = list(dict.fromkeys(m for *_, m in found if m))
    if len(monies) > 1:
        contract = [m for m in monies if "도급" in m and "하도급" not in m]
        chosen_money = contract[0] if contract else monies[0]
        found = [f for f in found if f[3] == chosen_money]

    groups = {g for _, _, g, _ in found if g}
    if len(groups) > 1:
        chosen = _pick(groups, "변경후")
        cols = [c for c, b, g, _ in found if g == chosen and not _is_total(b)]
        if not _has_qty(grid.ws, cols, data_row):
            chosen = _pick(groups, "변경전")
        found = [f for f in found if f[2] == chosen]
    return [(c, b) for c, b, _, _ in found if not _is_total(b)]


def _heading(text: str) -> tuple[tuple[str, ...] | None, str]:
    """"1-1 0102. 철근콘크리트공사" → (('1','1','01','02'), '철근콘크리트공사'). 번호 없으면 (None, 원문)."""
    m = _HEADING.match(text)
    if not m:
        return None, text.strip()
    head, tail, title = m.groups()
    code = head.split("-")
    if tail:
        code += [tail[i:i + 2] for i in range(0, len(tail), 2)]
    return tuple(code), title


class _Sections:
    """구분 행 번호의 앞부분 일치로 계층을 쌓는다."""

    def __init__(self):
        self.stack: list[tuple[tuple[str, ...], str]] = []

    def push(self, text: str) -> None:
        code, title = _heading(text)
        title = title.strip("= ").strip()
        if code is None:
            # 번호 없는 제목("=== CIP공사 ===")끼리는 형제: 가장 가까운 번호 제목의 자식으로 바꿔 끼운다.
            while self.stack and self.stack[-1][0][-1:] == ("*",):
                self.stack.pop()
            code = (self.stack[-1][0] if self.stack else ()) + ("*",)
            self.stack.append((code, title))
            return
        while self.stack:
            top = self.stack[-1][0]
            if len(top) < len(code) and code[:len(top)] == top:
                break
            self.stack.pop()
        self.stack.append((code, title))

    @property
    def path(self) -> str:
        return SECTION_SEP.join(t for _, t in self.stack if t)


def read_boq(path: str | Path) -> list[BoqLine]:
    path = Path(path)
    if path.name.startswith("~$"):
        raise ValueError(f"오피스 잠금 파일은 읽지 않음: {path.name}")
    wb = openpyxl.load_workbook(path, data_only=True)
    lines: list[BoqLine] = []
    for ws in wb.worksheets:
        for pattern, supply in _SHEETS:
            m = pattern.match(ws.title)
            if m:
                break
        else:
            continue
        grid = _Grid(ws)
        header = _find_header(grid)
        if header is None:
            continue
        qty_row, header_end, qty_cols, cols = header
        blocks = _qty_columns(grid, qty_row, qty_cols, header_end + 1)
        sections = _Sections()
        for r in range(header_end + 1, ws.max_row + 1):
            name = ws.cell(r, cols["품명"]).value
            if _is_marker_row(ws, r):
                if "구분" in _norm(ws.cell(r, 1).value) and name is not None and str(name).strip():
                    sections.push(str(name))
                continue
            if name is None or not str(name).strip() or _norm(name) in _TOTAL_NAMES:
                continue
            spec = ws.cell(r, cols["규격"]).value if "규격" in cols else None
            unit = ws.cell(r, cols["단위"]).value if "단위" in cols else None
            for c, block in blocks:
                qty = _qty(ws.cell(r, c).value)
                if not qty:
                    continue
                lines.append(BoqLine(
                    discipline=_DISCIPLINE[m.group(1)],
                    sheet=ws.title,
                    row=r,
                    name=str(name).strip(),
                    spec="" if spec is None else str(spec).strip(),
                    unit="" if unit is None else str(unit).strip(),
                    qty=qty,
                    block=block,
                    supply=supply,
                    section=sections.path,
                ))
    return _assign_sheet_blocks(wb, lines)


def _assign_sheet_blocks(wb, lines: list[BoqLine]) -> list[BoqLine]:
    """블록 열이 없는 블록별 시트(예: '내(건)3', '내역(건)<블록>')는 시트 이름 꼬리를 블록으로 쓴다.
    같은 파일의 블록 열 이름 중 꼬리를 포함하는 것이 하나뿐이면 그 이름으로 맞춘다(지급·사급 블록명 통일)."""
    from dataclasses import replace
    labels = {ln.block for ln in lines if ln.block}
    tails: dict[str, str] = {}
    for ws in wb.worksheets:
        for pattern, _ in _SHEETS:
            m = pattern.match(ws.title)
            if m:
                tail = ws.title[m.end():].strip(" _-()")
                if tail:
                    hit = [lb for lb in labels if tail in lb]
                    tails[ws.title] = hit[0] if len(hit) == 1 else tail
                break
    return [replace(ln, block=tails[ln.sheet]) if not ln.block and ln.sheet in tails else ln for ln in lines]

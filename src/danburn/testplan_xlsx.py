"""품질시험계획서 단독본(엑셀 .xlsx) — openpyxl 로 직접 만든다(오피스 프로그램 없음, CLAUDE.md 규칙 5).

hwpx 단독본(testplan)과 같은 rows·project 를 쓰고, 8.11 표의 칸 정의(HEAD·HEAD_CELLS·merge_plan·cell_texts)를
hwpx_out 에서 그대로 가져온다 → 두 형식의 칸·병합이 어긋나지 않는다.

시트: 표지 / 1.공사개요 / 2.시험계획-<분야>(분야마다) / 3.시험장비·교정 / 3.시험실 / 4.배치계획 (+ 미작성 자재).
시험계획 시트는 A4 가로·머리행 반복·한 쪽 너비 맞춤·공종 경계 쪽나눔. 숫자(계획물량·현장·의뢰 횟수)는 숫자형,
수식은 넣지 않는다(값 고정 — 받는 쪽 엑셀·뷰어 어디서나 같은 값). 글꼴은 이름만 적고 파일은 넣지 않는다.
원자적 저장(임시 파일 → 교체)은 호출하는 쪽(cli)이 한다: build_workbook 은 Workbook 을, save 는 경로에 쓴다.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

from . import plan_doc
from .hwpx_out import (ALIGN, C_EXT, C_QTY, C_SITE, HEAD, HEAD_CELLS, UNLISTED_HEAD, WIDTHS_MM,
                       _unlisted_texts, cell_texts, compact_date, merge_plan)
from .model import PlanRow
from .testplan import (CALIB_COLS, DOC_TITLE, EQUIP_COLS, PARTS, PARTS_BASIS, SCAN, STAFF_COLS, STAFF_PROOFS, TODO,
                       overview_values, room_text, staff_rows, table_rows)

FONT = "맑은 고딕"                    # 이름만(파일 동봉 없음). 없는 PC 는 엑셀이 대체 글꼴로 연다
PT_TITLE, PT_HEAD, PT_BODY, PT_NOTE = 14, 9, 9, 8
HEAD_FILL = PatternFill("solid", fgColor="E8E8E8")
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PLAN_TITLE_ROW, PLAN_HEAD_ROW = 1, 2          # 시험계획 시트: 1행 제목, 2~3행 머리(쪽마다 반복), 4행부터 본문
PLAN_BODY_ROW = PLAN_HEAD_ROW + 2
LAND_BODY_MM = 297 - 2 * 10                    # 가로 A4, 여백 10mm 기준 본문폭
MM_PER_CHAR = 1.9                              # 엑셀 열 너비 1(기본 글꼴 글자 하나) ≈ 1.9mm
QTY_FORMAT = "#,##0.###"
MARGIN_IN = 10 / 25.4


def _font(pt: float = PT_BODY, bold: bool = False) -> Font:
    return Font(name=FONT, size=pt, bold=bold)


def _plan_widths() -> list[float]:
    """8.11 칸 너비(hwpx 세로 A4 mm 비율 그대로)를 가로 A4 본문폭에 맞춘 엑셀 열 너비."""
    return [round(x * LAND_BODY_MM / sum(WIDTHS_MM) / MM_PER_CHAR, 1) for x in WIDTHS_MM]


def _page(ws, *, landscape: bool, title_rows: str | None = None, area: str | None = None):
    """A4, 한 쪽 너비 맞춤(높이는 쪽 수 제한 없음), 여백 10mm, 가운데 정렬 없음."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    m = ws.page_margins
    m.left = m.right = m.top = m.bottom = MARGIN_IN
    m.header = m.footer = MARGIN_IN / 2
    ws.oddFooter.center.text = "&P / &N"          # 쪽 번호(바닥글)
    ws.oddFooter.center.size = PT_NOTE
    if title_rows:
        ws.print_title_rows = title_rows
    if area:
        ws.print_area = area


def _cell(ws, r: int, c: int, value, *, bold=False, pt=PT_BODY, align="left", head=False, border=True, wrap=True,
          fmt: str | None = None):
    cell = ws.cell(row=r, column=c, value=value)
    cell.font = _font(PT_HEAD if head else pt, bold or head)
    cell.alignment = Alignment(horizontal="center" if head else align, vertical="center", wrap_text=wrap)
    if border:
        cell.border = BOX
    if head:
        cell.fill = HEAD_FILL
    if fmt:
        cell.number_format = fmt
    return cell


def _box_merge(ws, r0: int, c0: int, r1: int, c1: int):
    """병합 + 병합 범위 모든 칸에 테두리(엑셀은 병합 범위 칸마다 선을 따로 가진다)."""
    if (r0, c0) != (r1, c1):
        ws.merge_cells(start_row=r0, start_column=c0, end_row=r1, end_column=c1)
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            ws.cell(row=r, column=c).border = BOX


def _table(ws, top: int, head: list[str], rows: list[list], *, left_cols=(), widths: list[float] | None = None) -> int:
    """머리 1행 + 자료 행. 반환: 표 다음 빈 행 번호."""
    for c, h in enumerate(head, 1):
        _cell(ws, top, c, h, head=True)
    for i, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            _cell(ws, top + i, c, v, align="left" if c - 1 in left_cols else "center")
    if widths:
        for c, wd in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = wd
    return top + len(rows) + 2


def _title(ws, r: int, text: str, ncol: int, pt=PT_TITLE):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    c = ws.cell(row=r, column=1, value=text)
    c.font = _font(pt, True)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[r].height = pt * 1.8


def _note(ws, r: int, text: str, ncol: int):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    c = ws.cell(row=r, column=1, value=text)
    c.font = _font(PT_NOTE)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _scan(ws, r: int, label: str, ncol: int, height_rows: int, note: str = "") -> int:
    """(스캔첨부) 빈 칸: 여러 행을 병합한 큰 칸. 반환: 다음 빈 행."""
    _box_merge(ws, r, 1, r + height_rows - 1, ncol)
    c = ws.cell(row=r, column=1, value=f"{SCAN} {label}" + (f"\n{note}" if note else ""))
    c.font = _font(12, True)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return r + height_rows + 1


# ── 시트들 ──────────────────────────────────────────────────────────────

def _sheet_cover(wb: Workbook, project: dict, revs: list[dict]):
    ws = wb.active
    ws.title = "표지"
    ncol = 3
    for col, wd in zip("ABC", (10, 30, 50)):
        ws.column_dimensions[col].width = wd
    _title(ws, 1, _plain(project.get("공사명")) or TODO, ncol, 14)
    _title(ws, 2, DOC_TITLE, ncol, 22)
    cur = revs[-1]
    info = [("문서번호", _plain(project.get("문서번호"))), ("개정번호", f"Rev.{cur['개정']}"),
            ("개정일자", compact_date(cur["일자"])), ("작성", _plain(project.get("회사명")) or _plain(project.get("시공자")))]
    r = 4
    for k, v in info:
        _cell(ws, r, 1, k, bold=True, align="center")
        _box_merge(ws, r, 2, r, ncol)
        _cell(ws, r, 2, v)
        r += 1
    r += 1
    _cell(ws, r, 1, "목  차", bold=True, border=False)
    r = _table(ws, r + 1, ["번호", "제목", "내용"], [[no, t, "\n".join(s)] for no, t, s in PARTS], left_cols={1, 2})
    _note(ws, r - 1, f"구성 근거: {PARTS_BASIS}", ncol)
    _cell(ws, r + 1, 1, "개정 이력", bold=True, border=False)
    _table(ws, r + 2, ["개정번호", "개정일자", "개정 사유"],
           [[int(x["개정"]), compact_date(x["일자"]), _plain(x.get("사유"))] for x in revs], left_cols={2})
    _page(ws, landscape=False)


def _sheet_overview(wb: Workbook, project: dict):
    ws = wb.create_sheet("1.공사개요")
    ncol = 4
    for col, wd in zip("ABCD", (18, 22, 22, 40)):
        ws.column_dimensions[col].width = wd
    _title(ws, 1, "1. 개요 — 가. 공사개요", ncol)
    vals = overview_values(project)
    r = 3
    for k in plan_doc.OVERVIEW_ORDER:
        v = vals.get(k)
        if v is None or isinstance(v, (list, dict)) or not str(v).strip():
            continue
        _cell(ws, r, 1, plan_doc.OVERVIEW_LABEL.get(k, k), bold=True, align="center")
        _box_merge(ws, r, 2, r, ncol)
        _cell(ws, r, 2, str(v))
        r += 1
    r += 1
    _title(ws, r, "나. 공사수행 조직", ncol, 11)
    people = [x for x in (project.get("조직") or []) if isinstance(x, dict) and _plain(x.get("성명"))]
    rows = [[_plain(x.get("직무")), _plain(x.get("성명")), _plain(x.get("자격")), _plain(x.get("담당"))] for x in people]
    r = _table(ws, r + 1, ["직무", "성명", "자격", "담당업무"], rows or [[""] * 4 for _ in range(5)], left_cols={3})
    _note(ws, r - 1, "조직도 그림은 한글 단독본(품질시험계획서.hwpx) 1. 개요에 있다.", ncol)
    _page(ws, landscape=False)


def _plan_sheet(wb: Workbook, disc: str, rows: list[PlanRow], *, basis_version: str, generated_note: str):
    """분야 하나의 시험계획 시트. 칸·병합은 hwpx 8.11 표와 같다(merge_plan)."""
    ws = wb.create_sheet(_sheet_name(f"2.시험계획-{disc}"))
    ncol = len(HEAD)
    for c, wd in enumerate(_plan_widths(), 1):
        ws.column_dimensions[get_column_letter(c)].width = wd
    _title(ws, PLAN_TITLE_ROW, f"2. 품질시험 및 검사계획 — {disc}", ncol)
    for r0, c0, rs, cs, text in HEAD_CELLS:
        r, c = PLAN_HEAD_ROW + r0, c0 + 1
        _cell(ws, r, c, text, head=True)
        _box_merge(ws, r, c, r + rs - 1, c + cs - 1)
        for rr in range(r, r + rs):
            for cc in range(c, c + cs):
                ws.cell(row=rr, column=cc).fill = HEAD_FILL
    texts = [cell_texts(x) for x in rows]
    plan = merge_plan(rows)
    for i, (row, t) in enumerate(zip(rows, texts)):
        r = PLAN_BODY_ROW + i
        for c in range(ncol):
            value, fmt = t[c], None
            if c == C_QTY and row.qty:
                value, fmt = row.qty, QTY_FORMAT            # 숫자형(엑셀에서 합계·정렬 가능)
            elif c == C_SITE:
                value = int(row.count_site or 0)
            elif c == C_EXT:
                value = int(row.count_external or 0)
            _cell(ws, r, c + 1, value, align=ALIGN[c].lower().replace("justify", "left"), fmt=fmt)
    for c, runs in plan.items():
        for s, e in runs:
            if e > s:
                _box_merge(ws, PLAN_BODY_ROW + s, c + 1, PLAN_BODY_ROW + e, c + 1)
    last = PLAN_BODY_ROW + len(rows) - 1
    # 공종이 바뀌는 곳에서 쪽을 나눈다(공종 첫 행 앞) — 쪽마다 공종 이름이 보이게
    for s, _e in plan[0][1:]:
        ws.row_breaks.append(Break(id=PLAN_BODY_ROW + s - 1))
    r = last + 2
    for text in (f"근거 기준: {basis_version}" if basis_version else "", generated_note):
        if text:
            _note(ws, r, text, ncol)
            r += 1
    ws.freeze_panes = ws.cell(row=PLAN_BODY_ROW, column=1)
    _page(ws, landscape=True, title_rows=f"{PLAN_HEAD_ROW}:{PLAN_HEAD_ROW + 1}",
          area=f"A1:{get_column_letter(ncol)}{max(last, PLAN_BODY_ROW)}")
    return ws


def _sheet_unlisted(wb: Workbook, unlisted: list[dict]):
    ws = wb.create_sheet("2.미작성 자재(확인)")
    ncol = len(UNLISTED_HEAD)
    _title(ws, 1, "시험계획 미작성 자재 (확인 필요)", ncol, 12)
    rows = []
    for it in unlisted:
        t = _unlisted_texts(it)
        lines = it.get("lines")
        rows.append([t[0], t[1], t[2], lines if isinstance(lines, int) else "-", t[4]])
    _table(ws, 3, UNLISTED_HEAD, rows, left_cols={1, 2, 4}, widths=[14, 24, 24, 10, 40])
    _page(ws, landscape=False, title_rows="3:3")


def _sheet_equipment(wb: Workbook, project: dict):
    ws = wb.create_sheet("3.시험장비·교정")
    ncol = len(CALIB_COLS) + 1
    _title(ws, 1, "3. 품질시험 시설 — 가. 시험·검사 장비 / 나. 교정계획", ncol)
    cols = ["번호"] + list(dict.fromkeys(EQUIP_COLS[:-1] + CALIB_COLS))          # 두 표를 한 표로(같은 장비 목록)
    data = table_rows(project, "시험장비", cols[1:])
    rows = [[i if any(r) else "", *r] for i, r in enumerate(data, 1)]
    r = _table(ws, 3, cols, rows, left_cols={1, 2, len(cols) - 1},
               widths=[6, 22, 18, 7, 7, 14, 14, 10, 20][:len(cols)])
    _note(ws, r - 1, "교정은 정해진 주기 또는 사용 전에 받고 교정성적서를 보관한다(품질관리계획서 7.2).", len(cols))
    _page(ws, landscape=False, title_rows="3:3")


def _sheet_room(wb: Workbook, project: dict):
    ws = wb.create_sheet("3.시험실")
    ncol = 4
    for col, wd in zip("ABCD", (20, 25, 25, 25)):
        ws.column_dimensions[col].width = wd
    _title(ws, 1, "3. 품질시험 시설 — 다. 시험실", ncol)
    for r, (k, v) in enumerate((("시험실 규모(기준)", room_text(project)),
                                ("시험실 면적(실제)", "가로      m × 세로      m =        ㎡")), 3):
        _cell(ws, r, 1, k, bold=True, align="center")
        _box_merge(ws, r, 2, r, ncol)
        _cell(ws, r, 2, v)
    _scan(ws, 6, "시험실 배치평면도", ncol, 30, "시험실·양생 수조·시험기구 배치와 치수를 표시한 평면도를 붙인다")
    _page(ws, landscape=False)


def _sheet_staff(wb: Workbook, project: dict):
    ws = wb.create_sheet("4.배치계획")
    ncol = len(STAFF_COLS)
    _title(ws, 1, "4. 품질관리자 배치계획", ncol)
    r = _table(ws, 3, STAFF_COLS, staff_rows(project), left_cols={3, 4}, widths=[16, 16, 12, 24, 24])
    _note(ws, r - 1, "등급·배치 기준: 「건설기술 진흥법 시행규칙」 제50조제4항·별표5.", ncol)
    r += 1
    _title(ws, r, "증빙(사람마다 붙인다 — 주민등록번호 뒷자리 등 개인정보는 가린 사본)", ncol, 10)
    r += 1
    for proof in STAFF_PROOFS:
        r = _scan(ws, r, proof, ncol, 4)
    _page(ws, landscape=False)


# ── 공개 함수 ────────────────────────────────────────────────────────────

def _plain(v) -> str:
    return "" if v is None else str(v).strip()


def _sheet_name(name: str) -> str:
    """엑셀 시트 이름 규칙(31자, []:*?/\\ 금지)."""
    for ch in "[]:*?/\\":
        name = name.replace(ch, "·")
    return name[:31]


def build_workbook(rows: list[PlanRow], project: dict, *, basis_version: str = "", revision: int | str = 0,
                   date: str = "", unlisted: list[dict] | None = None, generated_note: str = "") -> Workbook:
    """품질시험계획서 통합문서. 개정이력 검사는 hwpx 단독본과 같다(어긋나면 ValueError)."""
    project = dict(project or {})
    revs = plan_doc._revisions(project, revision, date or _plain(project.get("제정일자")))
    wb = Workbook()
    _sheet_cover(wb, project, revs)
    _sheet_overview(wb, project)
    for disc in dict.fromkeys(r.discipline for r in rows):
        _plan_sheet(wb, disc, [r for r in rows if r.discipline == disc],
                    basis_version=basis_version, generated_note=generated_note)
    if unlisted:
        _sheet_unlisted(wb, unlisted)
    _sheet_equipment(wb, project)
    _sheet_room(wb, project)
    _sheet_staff(wb, project)
    return wb


def save_workbook(rows: list[PlanRow], project: dict, out: str | Path, **kw) -> Path:
    """build_workbook 결과를 out 에 쓴다(원자적 교체는 호출하는 쪽 — cli 의 _tmp_for/_commit)."""
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(rows, project, **kw).save(out)
    return out


__all__ = ["build_workbook", "save_workbook"]

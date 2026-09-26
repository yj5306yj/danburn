#!/usr/bin/env python3
"""공개용 예제: 완전 합성 공동주택 도급내역서(xlsx)를 만든다.

실제 현장·회사·수량과 무관한 합성 자료다. 실제 도급내역서의 모양만 흉내 낸다.
- 시트: 지급(건)·지급(토)(지급자재), 내역(건)·내역(토)·내역(기)(사급 내역), 그리고 읽지 않는 원가계산서
- 머리 두 줄: 블록 열 "가동"·"나동"·"합계" 아래에 "수량"·"금액"
- 1열 표시: 목차·구분(번호 제목)·소계 행. 번호 제목은 `1-1. 아파트 > 1-1 01. 상부공사 > 1-1 0101. …` 계층

사용: scripts/make_example_boq.py [--out examples/out/example_boq.xlsx]
"""
from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

BLOCKS = ("가동", "나동")
DEFAULT_OUT = Path(__file__).resolve().parents[1] / "examples" / "out" / "example_boq.xlsx"

# 시트 → [(구분 행 제목, [(품명, 규격, 단위, 가동 수량, 나동 수량), …] 또는 None)]
# 품목이 None 인 제목은 하위 구분을 여는 머리다. 수량은 모두 합성 값이다.
SHEETS: dict[str, list[tuple[str, list[tuple] | None]]] = {
    "지급(건)": [
        ("1-1. 아파트", None),
        ("1-1 01. 상부공사", None),
        ("1-1 0101. 철근콘크리트공사", [
            ("레미콘", "25-24-150", "M3", 2400, 2000),
            ("레미콘", "25-27-150", "M3", 850, 700),
            ("레미콘", "25-30-150", "M3", 400, 350),
        ]),
        ("1-1 02. 지하주차장", None),
        ("1-1 0201. 철근콘크리트공사", [
            ("레미콘", "25-24-150", "M3", 1100, 950),
            ("레미콘", "25-18-80", "M3", 90, 80),
        ]),
    ],
    "지급(토)": [
        ("2-1. 부대토목", None),
        ("2-1 01. 우수공사", [
            ("레미콘", "25-18-80", "M3", 60, 45),
        ]),
        ("2-1 02. 포장공사", [
            ("레미콘", "25-21-120", "M3", 40, 30),
        ]),
    ],
    "내역(건)": [
        ("1-1. 아파트", None),
        ("1-1 01. 상부공사", None),
        ("1-1 0101. 철근콘크리트공사", [
            ("이형봉강(SD500)", "H-13", "TON", 180.5, 150.2),
            ("이형봉강(SD500)", "H-16", "TON", 95.0, 80.4),
            ("이형봉강(SD400)", "H-10", "TON", 120.0, 100.6),
            ("철근 가공조립", "보통", "TON", 395.5, 331.2),
            ("거푸집용합판", "12mm 코팅", "M2", 5200, 4400),
            ("PVC지수판 설치", "200×5mm", "M", 320, 280),
            ("보통인부", "", "인", 1500, 1300),
        ]),
        ("1-1 0102. 조적공사", [
            ("시멘트벽돌", "190×90×57", "천매", 85.0, 72.0),
            ("콘크리트블록", "390×190×100", "매", 4200, 3600),
            ("벽돌 쌓기", "1.0B", "천매", 85.0, 72.0),
            ("건조시멘트모르타르", "조적용 40kg", "포", 3100, 2600),
            ("포틀랜드시멘트", "1종 40kg", "포", 900, 760),
        ]),
        ("1-1 0103. 방수공사", [
            ("우레탄도막방수", "노출형 3mm", "M2", 1300, 1100),
            ("시멘트액체방수", "바닥 2차", "M2", 2100, 1800),
        ]),
        ("1-1 0104. 단열공사", [
            ("비드법보온판", "2종1호 100mm", "M2", 3400, 2900),
            ("압출법단열재", "특호 50mm", "M2", 1200, 1000),
        ]),
        ("1-1 0105. 창호 및 유리공사", [
            ("합성수지창호", "WW1 1,200×1,500", "개소", 240, 200),
            ("합성수지창호", "WW2 2,400×2,300", "개소", 120, 100),
            ("복층유리", "24mm(6+12+6)", "M2", 1500, 1250),
            ("강화유리", "12mm", "M2", 180, 150),
        ]),
        ("1-1 0106. 타일공사", [
            ("자기질타일", "300×300", "M2", 2600, 2200),
            ("도기질타일", "250×400", "M2", 3800, 3200),
            ("압착시멘트", "40kg", "포", 1400, 1200),
        ]),
        ("1-1 0107. 도장공사", [
            ("수성페인트", "1급 내부", "M2", 9000, 7600),
            ("조합페인트", "1급", "M2", 1100, 950),
        ]),
        ("1-1 0108. 수장공사", [
            ("석고보드", "9.5mm", "M2", 6500, 5500),
        ]),
        ("1-1 0109. 기타공사", [
            ("실링재", "실리콘계", "M", 1250, 1050),
            ("주방가구", "84형", "SET", 120, 100),
        ]),
    ],
    "내역(토)": [
        ("2-1. 부대토목", None),
        ("2-1 01. 토공사", [
            ("터파기", "토사", "M3", 9000, 7500),
            ("되메우기", "장비+인력", "M3", 3200, 2700),
            ("동상방지층", "모래자갈", "M3", 700, 600),
        ]),
        ("2-1 02. 포장공사", [
            ("아스콘", "표층 #78", "TON", 260, 220),
            ("아스팔트포장", "표층 t=50mm", "M2", 2300, 1900),
            ("경계블록", "180×205×1000", "M", 520, 430),
            ("보도블록", "인터로킹 60mm", "M2", 900, 750),
        ]),
        ("2-1 03. 우수공사", [
            ("흄관", "D300", "M", 240, 200),
            ("토목용부직포", "4.0mm", "M2", 1100, 900),
        ]),
    ],
    "내역(기)": [
        ("3-1. 기계설비", None),
        ("3-1 01. 급배수위생설비", [
            ("PVC관", "VG1 50mm", "M", 800, 680),
            ("PVC관", "VG1 75mm", "M", 520, 440),
            ("PVC관", "VG1 100mm", "M", 610, 520),
            ("PVC관", "VG2 150mm", "M", 150, 120),
            ("그라스울 보온통", "25T", "M", 900, 760),
            ("배관공", "", "인", 300, 250),
        ]),
    ],
}


def _sheet(wb: openpyxl.Workbook, title: str, sections: list) -> None:
    ws = wb.create_sheet(title)
    ws.append([f"도 급 내 역 서 — {title} (합성 예제, 실제 현장과 무관)"])
    ws.append(["구분", "품명", "규격", "단위", BLOCKS[0], None, BLOCKS[1], None, "합 계", None])
    ws.append([None, None, None, None, "수량", "금액", "수량", "금액", "수량", "금액"])
    for col in ("A", "B", "C", "D"):
        ws.merge_cells(f"{col}2:{col}3")
    for left, right in (("E", "F"), ("G", "H"), ("I", "J")):
        ws.merge_cells(f"{left}2:{right}2")
    ws.append(["목차", sections[0][0]])
    for heading, items in sections:
        ws.append(["구분", heading])
        if not items:
            continue
        for name, spec, unit, qa, qb in items:
            ws.append([None, name, spec, unit, qa, None, qb, None, round(qa + qb, 3), None])
        ws.append(["소계", None, None, None, None, None, None, None, None, None])   # 금액 소계 자리(합성 예제는 금액을 비워 둔다)
    for cell in ws[2] + ws[3]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col, width in zip("ABCDEFGHIJ", (8, 24, 22, 6, 10, 10, 10, 10, 10, 10)):
        ws.column_dimensions[col].width = width


def build(out: Path) -> Path:
    wb = openpyxl.Workbook()
    cost = wb.active
    cost.title = "원가계산서"                        # 읽지 않는 시트(danburn 은 조용히 건너뜀)
    cost.append(["원 가 계 산 서 (합성 예제)"])
    cost.append(["비목", "금액"])
    for title, sections in SHEETS.items():
        _sheet(wb, title, sections)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="합성 공동주택 도급내역서 예제(xlsx)를 만든다")
    ap.add_argument("--out", default=str(DEFAULT_OUT), help=f"출력 xlsx (기본 {DEFAULT_OUT.relative_to(DEFAULT_OUT.parents[2])})")
    a = ap.parse_args(argv)
    print(build(Path(a.out)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
